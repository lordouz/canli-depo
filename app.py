"""
PET Resin Komple ERP v2.9.0

Kurulum:
    pip install streamlit pandas openpyxl

Çalıştırma:
    streamlit run PET_Resin_ERP_v2_7.py

Veriler, aynı klasörde otomatik oluşturulan pet_resin_erp_v2_7.db dosyasında
kalıcı olarak saklanır.
"""

import io
import json
import os
import sqlite3
import tempfile
import uuid
from contextlib import contextmanager
from datetime import date, datetime
from pathlib import Path

import pandas as pd
import streamlit as st


APP_NAME = "PET Resin Komple ERP"
APP_VERSION = "v2.9.0"
DB_PATH = Path(
    os.getenv(
        "PET_ERP_DB_PATH",
        str(Path(__file__).resolve().with_name("pet_resin_erp_v2_7.db")),
    )
)

CATEGORIES = [
    "Hammadde",
    "Yardımcı Kimyasal",
    "Ambalaj",
    "Ara Mamul",
    "Mamul",
]
UNITS = ["Kg", "Adet", "Ton", "lt"]
ORDER_STATUS_LABELS = {
    "PLANNED": "Planlandı",
    "IN_PROGRESS": "Üretimde",
    "COMPLETED": "Tamamlandı",
    "CANCELLED": "İptal Edildi",
}
MOVEMENT_LABELS = {
    "RECEIPT": "Mal Kabul",
    "TRANSFER_OUT": "Transfer Çıkış",
    "TRANSFER_IN": "Transfer Giriş",
    "ADJUSTMENT_OUT": "Fire / Stok Düşümü",
    "ADJUSTMENT_IN": "Sayım Fazlası / Stok İlavesi",
    "PRODUCTION_CONSUMPTION": "Üretim Sarfiyatı",
    "PRODUCTION_OUTPUT": "Üretim Çıktısı",
    "PRODUCTION_WASTE_RECOVERY": "Üretim Ziyanı Geri Kazanım",
    "SHIPMENT": "Müşteri Sevkiyatı",
    "PRODUCTION_CANCEL": "Üretim İptal Ters Kaydı",
    "SHIPMENT_CANCEL": "Sevkiyat İptal Ters Kaydı",
}

WASTE_DISPOSITION_LABELS = {
    "NONE": "Ziyan Yok",
    "RECOVERABLE": "Tekrar Kullanılabilir (K2 Eritmelik)",
    "DISPOSAL": "Kullanılamaz / Bertaraf",
}

WASTE_TYPES = [
    "Başlangıç-Duruş Ziyanı",
    "Renk Geçiş Ziyanı",
    "Elek / Filtre Ziyanı",
    "Kalite Uygunsuzluğu",
    "Kontamine Ürün",
    "Dökülme",
    "Diğer",
]


class ERPError(Exception):
    """Kullanıcıya güvenli biçimde gösterilebilecek iş kuralı hatası."""


def now_text():
    return datetime.now().strftime("%Y-%m-%d %H:%M:%S")


def normalize_text(value):
    return " ".join(str(value or "").strip().split())


def normalize_code(value):
    return normalize_text(value).upper()


def make_reference(prefix):
    suffix = uuid.uuid4().hex[:5].upper()
    return f"{prefix}-{datetime.now().strftime('%Y%m%d-%H%M%S')}-{suffix}"


@contextmanager
def db_session(write=False):
    DB_PATH.parent.mkdir(parents=True, exist_ok=True)
    connection = sqlite3.connect(str(DB_PATH), timeout=30)
    connection.row_factory = sqlite3.Row
    connection.execute("PRAGMA foreign_keys = ON")
    connection.execute("PRAGMA busy_timeout = 30000")
    connection.execute("PRAGMA journal_mode = WAL")
    try:
        if write:
            connection.execute("BEGIN IMMEDIATE")
        yield connection
        if write and connection.in_transaction:
            connection.commit()
    except Exception:
        if connection.in_transaction:
            connection.rollback()
        raise
    finally:
        connection.close()


def query_df(sql, params=()):
    with db_session() as connection:
        return pd.read_sql_query(sql, connection, params=params)


def query_rows(sql, params=()):
    with db_session() as connection:
        return [dict(row) for row in connection.execute(sql, params).fetchall()]


def initialize_database():
    schema = """
    CREATE TABLE IF NOT EXISTS items (
        id INTEGER PRIMARY KEY AUTOINCREMENT,
        code TEXT NOT NULL UNIQUE COLLATE NOCASE,
        name TEXT NOT NULL UNIQUE COLLATE NOCASE,
        category TEXT NOT NULL,
        unit TEXT NOT NULL,
        min_stock REAL NOT NULL DEFAULT 0 CHECK (min_stock >= 0),
        active INTEGER NOT NULL DEFAULT 1 CHECK (active IN (0, 1)),
        created_at TEXT NOT NULL
    );

    CREATE TABLE IF NOT EXISTS warehouses (
        id INTEGER PRIMARY KEY AUTOINCREMENT,
        code TEXT NOT NULL UNIQUE COLLATE NOCASE,
        name TEXT NOT NULL UNIQUE COLLATE NOCASE,
        warehouse_type TEXT NOT NULL DEFAULT 'Genel',
        active INTEGER NOT NULL DEFAULT 1 CHECK (active IN (0, 1)),
        created_at TEXT NOT NULL
    );

    CREATE TABLE IF NOT EXISTS lots (
        id INTEGER PRIMARY KEY AUTOINCREMENT,
        item_id INTEGER NOT NULL,
        lot_no TEXT NOT NULL COLLATE NOCASE,
        created_at TEXT NOT NULL,
        expiry_date TEXT,
        UNIQUE (item_id, lot_no),
        FOREIGN KEY (item_id) REFERENCES items(id)
    );

    CREATE TABLE IF NOT EXISTS stock_movements (
        id INTEGER PRIMARY KEY AUTOINCREMENT,
        movement_time TEXT NOT NULL,
        item_id INTEGER NOT NULL,
        lot_id INTEGER NOT NULL,
        warehouse_id INTEGER NOT NULL,
        quantity_delta REAL NOT NULL CHECK (quantity_delta <> 0),
        movement_type TEXT NOT NULL,
        reference_type TEXT NOT NULL,
        reference_no TEXT NOT NULL,
        notes TEXT,
        created_by TEXT NOT NULL DEFAULT 'Sistem',
        reversal_of INTEGER UNIQUE,
        created_at TEXT NOT NULL,
        FOREIGN KEY (item_id) REFERENCES items(id),
        FOREIGN KEY (lot_id) REFERENCES lots(id),
        FOREIGN KEY (warehouse_id) REFERENCES warehouses(id),
        FOREIGN KEY (reversal_of) REFERENCES stock_movements(id)
    );

    CREATE TABLE IF NOT EXISTS recipes (
        id INTEGER PRIMARY KEY AUTOINCREMENT,
        name TEXT NOT NULL UNIQUE COLLATE NOCASE,
        recipe_type TEXT NOT NULL,
        product_item_id INTEGER NOT NULL,
        basis_quantity REAL NOT NULL CHECK (basis_quantity > 0),
        version INTEGER NOT NULL DEFAULT 1,
        active INTEGER NOT NULL DEFAULT 1 CHECK (active IN (0, 1)),
        created_at TEXT NOT NULL,
        updated_at TEXT NOT NULL,
        FOREIGN KEY (product_item_id) REFERENCES items(id)
    );

    CREATE TABLE IF NOT EXISTS recipe_lines (
        id INTEGER PRIMARY KEY AUTOINCREMENT,
        recipe_id INTEGER NOT NULL,
        item_id INTEGER NOT NULL,
        quantity REAL NOT NULL CHECK (quantity > 0),
        UNIQUE (recipe_id, item_id),
        FOREIGN KEY (recipe_id) REFERENCES recipes(id) ON DELETE CASCADE,
        FOREIGN KEY (item_id) REFERENCES items(id)
    );

    CREATE TABLE IF NOT EXISTS production_orders (
        id INTEGER PRIMARY KEY AUTOINCREMENT,
        order_no TEXT NOT NULL UNIQUE COLLATE NOCASE,
        recipe_id INTEGER NOT NULL,
        recipe_name_snapshot TEXT NOT NULL,
        recipe_version INTEGER NOT NULL,
        product_item_id INTEGER NOT NULL,
        product_lot_no TEXT NOT NULL,
        source_warehouse_id INTEGER NOT NULL,
        output_warehouse_id INTEGER NOT NULL,
        planned_quantity REAL NOT NULL CHECK (planned_quantity > 0),
        gross_quantity REAL,
        actual_quantity REAL,
        waste_quantity REAL NOT NULL DEFAULT 0 CHECK (waste_quantity >= 0),
        waste_disposition TEXT NOT NULL DEFAULT 'NONE',
        waste_type TEXT,
        waste_reason TEXT,
        waste_item_id INTEGER,
        waste_lot_no TEXT,
        waste_warehouse_id INTEGER,
        status TEXT NOT NULL DEFAULT 'PLANNED',
        planned_date TEXT NOT NULL,
        started_at TEXT,
        completed_at TEXT,
        cancelled_at TEXT,
        notes TEXT,
        created_at TEXT NOT NULL,
        FOREIGN KEY (recipe_id) REFERENCES recipes(id),
        FOREIGN KEY (product_item_id) REFERENCES items(id),
        FOREIGN KEY (source_warehouse_id) REFERENCES warehouses(id),
        FOREIGN KEY (output_warehouse_id) REFERENCES warehouses(id),
        FOREIGN KEY (waste_item_id) REFERENCES items(id),
        FOREIGN KEY (waste_warehouse_id) REFERENCES warehouses(id)
    );

    CREATE TABLE IF NOT EXISTS production_order_materials (
        id INTEGER PRIMARY KEY AUTOINCREMENT,
        production_order_id INTEGER NOT NULL,
        item_id INTEGER NOT NULL,
        planned_quantity REAL NOT NULL CHECK (planned_quantity >= 0),
        actual_quantity REAL,
        UNIQUE (production_order_id, item_id),
        FOREIGN KEY (production_order_id) REFERENCES production_orders(id) ON DELETE CASCADE,
        FOREIGN KEY (item_id) REFERENCES items(id)
    );

    CREATE TABLE IF NOT EXISTS shipments (
        id INTEGER PRIMARY KEY AUTOINCREMENT,
        shipment_no TEXT NOT NULL UNIQUE COLLATE NOCASE,
        shipment_time TEXT NOT NULL,
        customer_name TEXT NOT NULL,
        dispatch_note_no TEXT NOT NULL,
        plate TEXT,
        status TEXT NOT NULL DEFAULT 'COMPLETED',
        notes TEXT,
        cancelled_at TEXT,
        created_at TEXT NOT NULL
    );

    CREATE TABLE IF NOT EXISTS shipment_lines (
        id INTEGER PRIMARY KEY AUTOINCREMENT,
        shipment_id INTEGER NOT NULL,
        item_id INTEGER NOT NULL,
        lot_id INTEGER NOT NULL,
        warehouse_id INTEGER NOT NULL,
        quantity REAL NOT NULL CHECK (quantity > 0),
        movement_id INTEGER NOT NULL UNIQUE,
        FOREIGN KEY (shipment_id) REFERENCES shipments(id),
        FOREIGN KEY (item_id) REFERENCES items(id),
        FOREIGN KEY (lot_id) REFERENCES lots(id),
        FOREIGN KEY (warehouse_id) REFERENCES warehouses(id),
        FOREIGN KEY (movement_id) REFERENCES stock_movements(id)
    );

    CREATE INDEX IF NOT EXISTS idx_movements_balance
        ON stock_movements(item_id, lot_id, warehouse_id);
    CREATE INDEX IF NOT EXISTS idx_movements_date
        ON stock_movements(movement_time);
    CREATE INDEX IF NOT EXISTS idx_movements_reference
        ON stock_movements(reference_type, reference_no);
    CREATE INDEX IF NOT EXISTS idx_orders_status
        ON production_orders(status);
    """

    with db_session(write=True) as connection:
        connection.executescript(schema)
        migrate_production_waste_schema(connection)
        seed_master_data(connection)


def migrate_production_waste_schema(connection):
    """Eski veritabanlarına ziyan alanlarını veri kaybetmeden ekler."""
    existing_columns = {
        row["name"]
        for row in connection.execute("PRAGMA table_info(production_orders)").fetchall()
    }
    new_columns = {
        "gross_quantity": "REAL",
        "waste_quantity": "REAL NOT NULL DEFAULT 0",
        "waste_disposition": "TEXT NOT NULL DEFAULT 'NONE'",
        "waste_type": "TEXT",
        "waste_reason": "TEXT",
        "waste_item_id": "INTEGER REFERENCES items(id)",
        "waste_lot_no": "TEXT",
        "waste_warehouse_id": "INTEGER REFERENCES warehouses(id)",
    }
    for column_name, column_definition in new_columns.items():
        if column_name not in existing_columns:
            connection.execute(
                f"ALTER TABLE production_orders ADD COLUMN {column_name} {column_definition}"
            )

    connection.execute(
        """
        UPDATE production_orders
        SET gross_quantity = actual_quantity
        WHERE status = 'COMPLETED'
          AND gross_quantity IS NULL
          AND actual_quantity IS NOT NULL
        """
    )


def seed_master_data(connection):
    created_at = now_text()
    migrate_legacy_stock_cards(connection)
    seed_items = [
        ("HM-PTA", "PTA", "Hammadde", "Kg"),
        ("HM-MEG", "SAF MEG", "Hammadde", "Kg"),
        ("HM-IPA", "IPA", "Hammadde", "Kg"),
        ("HM-DEG", "DEG", "Hammadde", "Kg"),
        ("HM-KIRLIMEG", "KİRLİ MEG", "Hammadde", "Kg"),
        ("HM-YFLAKE", "YERLİ FLAKE", "Hammadde", "Kg"),
        ("HM-IFLAKE", "İTHAL FLAKE", "Hammadde", "Kg"),
        ("HM-K2", "K2 ERİTMELİK", "Hammadde", "Kg"),
        ("YK-SB2O3", "ANTİMON TRİOKSİT", "Yardımcı Kimyasal", "Kg"),
        ("YK-H3PO4", "FOSFORİK ASİT", "Yardımcı Kimyasal", "Kg"),
        ("YK-KIRMIZI", "KIRMIZI BOYA", "Yardımcı Kimyasal", "Kg"),
        ("YK-MAVI", "MAVİ BOYA", "Yardımcı Kimyasal", "Kg"),
        ("YK-REHEAT", "REHEAT", "Yardımcı Kimyasal", "Kg"),
        ("YK-TYZOR", "TYZOR AC 422", "Yardımcı Kimyasal", "Kg"),
        ("YK-TALK", "TALK", "Yardımcı Kimyasal", "Kg"),
        ("AMB-V1150", "1150 kg Virgin ürün BİG-BEG (Beyaz Kulak)", "Ambalaj", "Adet"),
        ("AMB-V1100", "1100 kg Virgin ürün BİG-BEG (Beyaz Kulak)", "Ambalaj", "Adet"),
        ("AMB-M1150", "1150 kg Baskısız MAVİ Kulak BİG-BEG", "Ambalaj", "Adet"),
        ("AMB-B1150", "1150 kg Baskısız BEYAZ Kulak BİG-BEG", "Ambalaj", "Adet"),
        ("AMB-R1150A", "1150 kg İç astarlı r-PET BİG-BEG (Yeşil Kulak)", "Ambalaj", "Adet"),
        ("AMB-R1100", "1100 kg r-PET BİG-BEG (Yeşil Kulak)", "Ambalaj", "Adet"),
        ("AMB-R1150", "1150 kg r-PET BİG-BEG (Yeşil Kulak)", "Ambalaj", "Adet"),
        ("AMB-V1100Y", "1100 kg Yeşil Kulak Virgin Baskılı BİG-BEG", "Ambalaj", "Adet"),
        ("AMB-YERLI-PAL", "YERLİ PALET", "Ambalaj", "Adet"),
        ("AMB-ITHAL-PAL", "İTHAL PALET", "Ambalaj", "Adet"),
        ("AMB-LINER", "LINER", "Ambalaj", "Adet"),
        ("AMB-SEP", "SEPERATÖR", "Ambalaj", "Adet"),
        ("AM-AMORF", "Standart Amorf Chips", "Ara Mamul", "Kg"),
    ]
    connection.executemany(
        """
        INSERT OR IGNORE INTO items
            (code, name, category, unit, min_stock, active, created_at)
        VALUES (?, ?, ?, ?, 0, 1, ?)
        """,
        [(*item, created_at) for item in seed_items],
    )
    remove_legacy_stock_cards(connection)

    seed_warehouses = [
        ("D01", "Depo 1", "Genel"),
        ("D02", "Depo 2", "Genel"),
        ("D03", "Depo 3", "Genel"),
    ]
    connection.executemany(
        """
        INSERT OR IGNORE INTO warehouses
            (code, name, warehouse_type, active, created_at)
        VALUES (?, ?, ?, 1, ?)
        """,
        [(*warehouse, created_at) for warehouse in seed_warehouses],
    )

    recipe_exists = connection.execute(
        "SELECT id FROM recipes WHERE name = ?",
        ("Standart Amorf Chips (Reaktör)",),
    ).fetchone()
    if not recipe_exists:
        product = connection.execute(
            "SELECT id FROM items WHERE name = ?", ("Standart Amorf Chips",)
        ).fetchone()
        if product:
            cursor = connection.execute(
                """
                INSERT INTO recipes
                    (name, recipe_type, product_item_id, basis_quantity,
                     version, active, created_at, updated_at)
                VALUES (?, 'Ara Mamul Reçetesi', ?, 1000, 1, 1, ?, ?)
                """,
                ("Standart Amorf Chips (Reaktör)", product["id"], created_at, created_at),
            )
            recipe_id = cursor.lastrowid
            default_lines = [
                ("PTA", 850.0),
                ("SAF MEG", 135.0),
                ("ANTİMON TRİOKSİT", 5.0),
            ]
            for item_name, quantity in default_lines:
                item_row = connection.execute(
                    "SELECT id FROM items WHERE name = ?", (item_name,)
                ).fetchone()
                if item_row:
                    connection.execute(
                        """
                        INSERT INTO recipe_lines (recipe_id, item_id, quantity)
                        VALUES (?, ?, ?)
                        """,
                        (recipe_id, item_row["id"], quantity),
                    )


def migrate_legacy_stock_cards(connection):
    migrations = [
        ("SAF DEG", "HM-DEG", "DEG", "Hammadde", "Kg"),
        ("Yurtiçi Standart Palet", "AMB-YERLI-PAL", "YERLİ PALET", "Ambalaj", "Adet"),
        ("Konteyner İhracat Paleti", "AMB-ITHAL-PAL", "İTHAL PALET", "Ambalaj", "Adet"),
        ("Karton Seperatör", "AMB-SEP", "SEPERATÖR", "Ambalaj", "Adet"),
        (
            "1100 Yeşil Kulak Virgin baskılı",
            "AMB-V1100Y",
            "1100 kg Yeşil Kulak Virgin Baskılı BİG-BEG",
            "Ambalaj",
            "Adet",
        ),
    ]
    for old_name, new_code, new_name, category, unit in migrations:
        old_item = connection.execute(
            "SELECT id FROM items WHERE name = ? COLLATE NOCASE", (old_name,)
        ).fetchone()
        if not old_item:
            continue
        conflict = connection.execute(
            """
            SELECT id FROM items
            WHERE id <> ? AND (code = ? COLLATE NOCASE OR name = ? COLLATE NOCASE)
            """,
            (old_item["id"], new_code, new_name),
        ).fetchone()
        if conflict:
            connection.execute("UPDATE items SET active = 0 WHERE id = ?", (old_item["id"],))
        else:
            connection.execute(
                """
                UPDATE items
                SET code = ?, name = ?, category = ?, unit = ?
                WHERE id = ?
                """,
                (new_code, new_name, category, unit, old_item["id"]),
            )


def remove_legacy_stock_cards(connection):
    legacy_names = ["HTM", "LPG", "MOTORİN", "MOTORIN", "DEŞE KIRIĞI"]
    for legacy_name in legacy_names:
        item = connection.execute(
            "SELECT id FROM items WHERE name = ? COLLATE NOCASE", (legacy_name,)
        ).fetchone()
        if not item:
            continue
        item_id = item["id"]
        reference_count = 0
        reference_queries = [
            "SELECT COUNT(*) AS count FROM lots WHERE item_id = ?",
            "SELECT COUNT(*) AS count FROM stock_movements WHERE item_id = ?",
            "SELECT COUNT(*) AS count FROM recipe_lines WHERE item_id = ?",
            "SELECT COUNT(*) AS count FROM recipes WHERE product_item_id = ?",
            "SELECT COUNT(*) AS count FROM production_orders WHERE product_item_id = ?",
            "SELECT COUNT(*) AS count FROM production_order_materials WHERE item_id = ?",
            "SELECT COUNT(*) AS count FROM shipment_lines WHERE item_id = ?",
        ]
        for sql in reference_queries:
            reference_count += int(connection.execute(sql, (item_id,)).fetchone()["count"])
        if reference_count:
            connection.execute("UPDATE items SET active = 0 WHERE id = ?", (item_id,))
        else:
            connection.execute("DELETE FROM items WHERE id = ?", (item_id,))


def get_or_create_lot(connection, item_id, lot_no, expiry_date=None):
    lot_no = normalize_code(lot_no)
    if not lot_no:
        raise ERPError("LOT / parti numarası boş bırakılamaz.")
    row = connection.execute(
        "SELECT id FROM lots WHERE item_id = ? AND lot_no = ?",
        (item_id, lot_no),
    ).fetchone()
    if row:
        return row["id"]
    cursor = connection.execute(
        """
        INSERT INTO lots (item_id, lot_no, created_at, expiry_date)
        VALUES (?, ?, ?, ?)
        """,
        (item_id, lot_no, now_text(), expiry_date),
    )
    return cursor.lastrowid


def get_balance(connection, item_id, lot_id, warehouse_id):
    row = connection.execute(
        """
        SELECT COALESCE(SUM(quantity_delta), 0) AS balance
        FROM stock_movements
        WHERE item_id = ? AND lot_id = ? AND warehouse_id = ?
        """,
        (item_id, lot_id, warehouse_id),
    ).fetchone()
    return float(row["balance"] or 0)


def add_movement(
    connection,
    *,
    item_id,
    lot_id,
    warehouse_id,
    quantity_delta,
    movement_type,
    reference_type,
    reference_no,
    movement_time=None,
    notes="",
    reversal_of=None,
):
    quantity_delta = float(quantity_delta)
    if abs(quantity_delta) < 0.0000001:
        raise ERPError("Sıfır miktarlı stok hareketi kaydedilemez.")
    cursor = connection.execute(
        """
        INSERT INTO stock_movements
            (movement_time, item_id, lot_id, warehouse_id, quantity_delta,
             movement_type, reference_type, reference_no, notes,
             created_by, reversal_of, created_at)
        VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, 'Sistem', ?, ?)
        """,
        (
            movement_time or now_text(),
            item_id,
            lot_id,
            warehouse_id,
            quantity_delta,
            movement_type,
            reference_type,
            normalize_code(reference_no),
            normalize_text(notes),
            reversal_of,
            now_text(),
        ),
    )
    return cursor.lastrowid


def stock_balance_df(positive_only=False, categories=None, warehouse_id=None):
    conditions = ["i.active = 1"]
    params = []
    if categories:
        placeholders = ",".join("?" for _ in categories)
        conditions.append(f"i.category IN ({placeholders})")
        params.extend(categories)
    if warehouse_id:
        conditions.append("w.id = ?")
        params.append(int(warehouse_id))
    where_sql = f"WHERE {' AND '.join(conditions)}" if conditions else ""
    having_sql = "HAVING SUM(sm.quantity_delta) > 0.000001" if positive_only else "HAVING ABS(SUM(sm.quantity_delta)) > 0.000001"
    sql = f"""
        SELECT
            i.id AS item_id,
            i.code AS "Stok Kodu",
            i.name AS "Malzeme / Ürün",
            i.category AS "Kategori",
            i.unit AS "Birim",
            l.id AS lot_id,
            l.lot_no AS "LOT No",
            w.id AS warehouse_id,
            w.name AS "Depo",
            ROUND(SUM(sm.quantity_delta), 6) AS "Güncel Stok"
        FROM stock_movements sm
        JOIN items i ON i.id = sm.item_id
        JOIN lots l ON l.id = sm.lot_id
        JOIN warehouses w ON w.id = sm.warehouse_id
        {where_sql}
        GROUP BY i.id, i.code, i.name, i.category, i.unit,
                 l.id, l.lot_no, w.id, w.name
        {having_sql}
        ORDER BY i.category, i.name, l.lot_no, w.name
    """
    return query_df(sql, tuple(params))


def movement_history_df(start_date=None, end_date=None, item_id=None):
    conditions = []
    params = []
    if start_date:
        conditions.append("date(sm.movement_time) >= date(?)")
        params.append(str(start_date))
    if end_date:
        conditions.append("date(sm.movement_time) <= date(?)")
        params.append(str(end_date))
    if item_id:
        conditions.append("sm.item_id = ?")
        params.append(int(item_id))
    where_sql = f"WHERE {' AND '.join(conditions)}" if conditions else ""
    df = query_df(
        f"""
        SELECT
            sm.id AS "Hareket ID",
            sm.movement_time AS "Tarih",
            i.code AS "Stok Kodu",
            i.name AS "Malzeme / Ürün",
            i.category AS "Kategori",
            l.lot_no AS "LOT No",
            w.name AS "Depo",
            sm.quantity_delta AS "Miktar",
            i.unit AS "Birim",
            sm.movement_type AS "Hareket Türü",
            sm.reference_no AS "Referans No",
            COALESCE(sm.notes, '') AS "Açıklama"
        FROM stock_movements sm
        JOIN items i ON i.id = sm.item_id
        JOIN lots l ON l.id = sm.lot_id
        JOIN warehouses w ON w.id = sm.warehouse_id
        {where_sql}
        ORDER BY sm.movement_time DESC, sm.id DESC
        """,
        tuple(params),
    )
    if not df.empty:
        df["Hareket Türü"] = df["Hareket Türü"].map(MOVEMENT_LABELS).fillna(df["Hareket Türü"])
    return df


def create_item(code, name, category, unit, min_stock):
    code = normalize_code(code)
    name = normalize_text(name)
    if not name:
        raise ERPError("Stok kartı adı zorunludur.")
    if category not in CATEGORIES or unit not in UNITS:
        raise ERPError("Kategori veya birim seçimi geçersiz.")
    if not code:
        prefixes = {
            "Hammadde": "HM",
            "Yardımcı Kimyasal": "YK",
            "Ambalaj": "AMB",
            "Ara Mamul": "AM",
            "Mamul": "M",
        }
        code = f"{prefixes[category]}-OTO-{uuid.uuid4().hex[:8].upper()}"
    with db_session(write=True) as connection:
        try:
            connection.execute(
                """
                INSERT INTO items
                    (code, name, category, unit, min_stock, active, created_at)
                VALUES (?, ?, ?, ?, ?, 1, ?)
                """,
                (code, name, category, unit, float(min_stock), now_text()),
            )
        except sqlite3.IntegrityError as exc:
            raise ERPError("Bu stok kodu veya stok kartı adı zaten kayıtlı.") from exc
    return code


def update_item(item_id, name, category, unit, min_stock, active):
    name = normalize_text(name)
    if not name:
        raise ERPError("Stok kartı adı boş bırakılamaz.")
    with db_session(write=True) as connection:
        current = connection.execute(
            "SELECT unit FROM items WHERE id = ?", (item_id,)
        ).fetchone()
        if not current:
            raise ERPError("Stok kartı bulunamadı.")
        if current["unit"] != unit:
            movement_count = connection.execute(
                "SELECT COUNT(*) AS count FROM stock_movements WHERE item_id = ?",
                (item_id,),
            ).fetchone()["count"]
            if movement_count:
                raise ERPError("Hareket görmüş bir stok kartının birimi değiştirilemez.")
        if not active:
            balance = connection.execute(
                """
                SELECT COALESCE(SUM(quantity_delta), 0) AS balance
                FROM stock_movements WHERE item_id = ?
                """,
                (item_id,),
            ).fetchone()["balance"]
            if abs(float(balance or 0)) > 0.000001:
                raise ERPError("Stoğu bulunan bir kart pasife alınamaz.")
            recipe_reference = connection.execute(
                """
                SELECT COUNT(*) AS count
                FROM recipes r
                LEFT JOIN recipe_lines rl ON rl.recipe_id = r.id
                WHERE r.active = 1
                  AND (r.product_item_id = ? OR rl.item_id = ?)
                """,
                (item_id, item_id),
            ).fetchone()["count"]
            if recipe_reference:
                raise ERPError("Aktif bir reçetede kullanılan stok kartı pasife alınamaz.")
            open_order_reference = connection.execute(
                """
                SELECT COUNT(*) AS count
                FROM production_orders po
                LEFT JOIN production_order_materials pom
                       ON pom.production_order_id = po.id
                WHERE po.status IN ('PLANNED', 'IN_PROGRESS')
                  AND (po.product_item_id = ? OR pom.item_id = ?)
                """,
                (item_id, item_id),
            ).fetchone()["count"]
            if open_order_reference:
                raise ERPError("Açık üretim emrinde kullanılan stok kartı pasife alınamaz.")
        try:
            connection.execute(
                """
                UPDATE items
                SET name = ?, category = ?, unit = ?, min_stock = ?, active = ?
                WHERE id = ?
                """,
                (name, category, unit, float(min_stock), int(bool(active)), item_id),
            )
        except sqlite3.IntegrityError as exc:
            raise ERPError("Bu stok kartı adı başka bir kartta kullanılıyor.") from exc


def create_warehouse(code, name, warehouse_type):
    code = normalize_code(code)
    name = normalize_text(name)
    if not code or not name:
        raise ERPError("Depo kodu ve depo adı zorunludur.")
    with db_session(write=True) as connection:
        try:
            connection.execute(
                """
                INSERT INTO warehouses
                    (code, name, warehouse_type, active, created_at)
                VALUES (?, ?, ?, 1, ?)
                """,
                (code, name, normalize_text(warehouse_type) or "Genel", now_text()),
            )
        except sqlite3.IntegrityError as exc:
            raise ERPError("Bu depo kodu veya depo adı zaten kayıtlı.") from exc


def set_warehouse_active(warehouse_id, active):
    with db_session(write=True) as connection:
        if not active:
            balance = connection.execute(
                """
                SELECT COALESCE(SUM(quantity_delta), 0) AS balance
                FROM stock_movements WHERE warehouse_id = ?
                """,
                (warehouse_id,),
            ).fetchone()["balance"]
            if abs(float(balance or 0)) > 0.000001:
                raise ERPError("İçinde stok bulunan depo pasife alınamaz.")
        connection.execute(
            "UPDATE warehouses SET active = ? WHERE id = ?",
            (int(bool(active)), warehouse_id),
        )


def create_receipt(item_id, warehouse_id, lot_no, quantity, movement_date, document_no, notes):
    quantity = float(quantity)
    if quantity <= 0:
        raise ERPError("Giriş miktarı sıfırdan büyük olmalıdır.")
    reference_no = normalize_code(document_no) or make_reference("GRS")
    with db_session(write=True) as connection:
        lot_id = get_or_create_lot(connection, item_id, lot_no)
        add_movement(
            connection,
            item_id=item_id,
            lot_id=lot_id,
            warehouse_id=warehouse_id,
            quantity_delta=quantity,
            movement_type="RECEIPT",
            reference_type="RECEIPT",
            reference_no=reference_no,
            movement_time=movement_date,
            notes=notes,
        )
    return reference_no


def transfer_stock(item_id, lot_id, source_warehouse_id, target_warehouse_id, quantity, notes):
    quantity = float(quantity)
    if source_warehouse_id == target_warehouse_id:
        raise ERPError("Kaynak ve hedef depo aynı olamaz.")
    if quantity <= 0:
        raise ERPError("Transfer miktarı sıfırdan büyük olmalıdır.")
    reference_no = make_reference("TRF")
    with db_session(write=True) as connection:
        balance = get_balance(connection, item_id, lot_id, source_warehouse_id)
        if quantity > balance + 0.000001:
            raise ERPError(f"Kaynak depoda yalnızca {balance:,.3f} kullanılabilir stok var.")
        add_movement(
            connection,
            item_id=item_id,
            lot_id=lot_id,
            warehouse_id=source_warehouse_id,
            quantity_delta=-quantity,
            movement_type="TRANSFER_OUT",
            reference_type="TRANSFER",
            reference_no=reference_no,
            notes=notes,
        )
        add_movement(
            connection,
            item_id=item_id,
            lot_id=lot_id,
            warehouse_id=target_warehouse_id,
            quantity_delta=quantity,
            movement_type="TRANSFER_IN",
            reference_type="TRANSFER",
            reference_no=reference_no,
            notes=notes,
        )
    return reference_no


def create_adjustment(
    *, item_id, warehouse_id, lot_no=None, lot_id=None, quantity, direction, reason
):
    quantity = float(quantity)
    if quantity <= 0:
        raise ERPError("Düzeltme miktarı sıfırdan büyük olmalıdır.")
    reference_no = make_reference("DUZ")
    with db_session(write=True) as connection:
        if direction == "OUT":
            if not lot_id:
                raise ERPError("Stok düşümü için lot seçilmelidir.")
            balance = get_balance(connection, item_id, lot_id, warehouse_id)
            if quantity > balance + 0.000001:
                raise ERPError(f"Seçilen stokta yalnızca {balance:,.3f} bulunuyor.")
            delta = -quantity
            movement_type = "ADJUSTMENT_OUT"
        else:
            lot_id = get_or_create_lot(connection, item_id, lot_no)
            delta = quantity
            movement_type = "ADJUSTMENT_IN"
        add_movement(
            connection,
            item_id=item_id,
            lot_id=lot_id,
            warehouse_id=warehouse_id,
            quantity_delta=delta,
            movement_type=movement_type,
            reference_type="ADJUSTMENT",
            reference_no=reference_no,
            notes=reason,
        )
    return reference_no


def resolve_or_create_recipe_product(connection, product_name, product_unit, recipe_type):
    product_name = normalize_text(product_name)
    if not product_name:
        raise ERPError("Üretilecek ürün adı boş bırakılamaz.")
    if product_unit not in UNITS:
        raise ERPError("Üretilecek ürün birimi geçersiz.")
    if recipe_type not in ["Ara Mamul Reçetesi", "Mamul Reçetesi"]:
        raise ERPError("Reçete sınıfı geçersiz.")

    target_category = "Ara Mamul" if recipe_type == "Ara Mamul Reçetesi" else "Mamul"
    existing = connection.execute(
        """
        SELECT id, code, name, category, unit, active
        FROM items WHERE name = ? COLLATE NOCASE
        """,
        (product_name,),
    ).fetchone()
    if existing:
        if existing["category"] != target_category:
            raise ERPError(
                f"'{existing['name']}' adı {existing['category']} kategorisinde kayıtlı. "
                f"Bu reçete için ürün kategorisi {target_category} olmalıdır."
            )
        if existing["unit"] != product_unit:
            raise ERPError(
                f"'{existing['name']}' stok kartının birimi {existing['unit']}. "
                f"Reçetede de aynı birimi seçin."
            )
        if not existing["active"]:
            connection.execute("UPDATE items SET active = 1 WHERE id = ?", (existing["id"],))
        return {
            "id": existing["id"],
            "code": existing["code"],
            "name": existing["name"],
            "created": False,
        }

    prefix = "AM" if target_category == "Ara Mamul" else "M"
    auto_code = f"{prefix}-OTO-{uuid.uuid4().hex[:8].upper()}"
    cursor = connection.execute(
        """
        INSERT INTO items
            (code, name, category, unit, min_stock, active, created_at)
        VALUES (?, ?, ?, ?, 0, 1, ?)
        """,
        (auto_code, product_name, target_category, product_unit, now_text()),
    )
    return {
        "id": cursor.lastrowid,
        "code": auto_code,
        "name": product_name,
        "created": True,
    }


def save_recipe(
    recipe_id,
    name,
    recipe_type,
    product_name,
    product_unit,
    basis_quantity,
    material_quantities,
):
    name = normalize_text(name)
    product_name = normalize_text(product_name)
    basis_quantity = float(basis_quantity)
    material_quantities = {
        int(item_id): float(quantity)
        for item_id, quantity in material_quantities.items()
        if float(quantity) > 0
    }
    if not name or not product_name or basis_quantity <= 0:
        raise ERPError("Reçete adı, üretilecek ürün adı ve baz üretim miktarı zorunludur.")
    if not material_quantities:
        raise ERPError("Reçeteye en az bir tüketim kalemi eklenmelidir.")
    if recipe_type not in ["Ara Mamul Reçetesi", "Mamul Reçetesi"]:
        raise ERPError("Reçete sınıfı geçersiz.")

    with db_session(write=True) as connection:
        try:
            product = resolve_or_create_recipe_product(
                connection, product_name, product_unit, recipe_type
            )
            product_item_id = product["id"]
            if int(product_item_id) in material_quantities:
                raise ERPError("Üretilecek ürün kendi reçetesinde tüketim kalemi olamaz.")
            if recipe_id:
                connection.execute(
                    """
                    UPDATE recipes
                    SET name = ?, recipe_type = ?, product_item_id = ?,
                        basis_quantity = ?, version = version + 1,
                        active = 1, updated_at = ?
                    WHERE id = ?
                    """,
                    (
                        name,
                        recipe_type,
                        product_item_id,
                        basis_quantity,
                        now_text(),
                        recipe_id,
                    ),
                )
                connection.execute("DELETE FROM recipe_lines WHERE recipe_id = ?", (recipe_id,))
            else:
                cursor = connection.execute(
                    """
                    INSERT INTO recipes
                        (name, recipe_type, product_item_id, basis_quantity,
                         version, active, created_at, updated_at)
                    VALUES (?, ?, ?, ?, 1, 1, ?, ?)
                    """,
                    (
                        name,
                        recipe_type,
                        product_item_id,
                        basis_quantity,
                        now_text(),
                        now_text(),
                    ),
                )
                recipe_id = cursor.lastrowid
            connection.executemany(
                """
                INSERT INTO recipe_lines (recipe_id, item_id, quantity)
                VALUES (?, ?, ?)
                """,
                [(recipe_id, item_id, quantity) for item_id, quantity in material_quantities.items()],
            )
        except sqlite3.IntegrityError as exc:
            raise ERPError("Bu reçete adı zaten kullanılıyor veya reçete bilgileri geçersiz.") from exc
    return {
        "recipe_id": recipe_id,
        "product_item_id": product_item_id,
        "product_name": product["name"],
        "product_code": product["code"],
        "product_created": product["created"],
    }


def create_production_order(
    order_no,
    recipe_id,
    planned_quantity,
    product_lot_no,
    source_warehouse_id,
    output_warehouse_id,
    planned_date,
    notes,
):
    order_no = normalize_code(order_no)
    product_lot_no = normalize_code(product_lot_no)
    planned_quantity = float(planned_quantity)
    if not order_no or not product_lot_no or planned_quantity <= 0:
        raise ERPError("Emir numarası, ürün lotu ve planlanan miktar zorunludur.")

    with db_session(write=True) as connection:
        recipe = connection.execute(
            """
            SELECT id, name, version, product_item_id, basis_quantity
            FROM recipes WHERE id = ? AND active = 1
            """,
            (recipe_id,),
        ).fetchone()
        if not recipe:
            raise ERPError("Aktif reçete bulunamadı.")
        lines = connection.execute(
            "SELECT item_id, quantity FROM recipe_lines WHERE recipe_id = ?",
            (recipe_id,),
        ).fetchall()
        if not lines:
            raise ERPError("Seçilen reçetede tüketim kalemi bulunmuyor.")
        try:
            cursor = connection.execute(
                """
                INSERT INTO production_orders
                    (order_no, recipe_id, recipe_name_snapshot, recipe_version,
                     product_item_id, product_lot_no, source_warehouse_id,
                     output_warehouse_id, planned_quantity, status,
                     planned_date, notes, created_at)
                VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, 'PLANNED', ?, ?, ?)
                """,
                (
                    order_no,
                    recipe_id,
                    recipe["name"],
                    recipe["version"],
                    recipe["product_item_id"],
                    product_lot_no,
                    source_warehouse_id,
                    output_warehouse_id,
                    planned_quantity,
                    str(planned_date),
                    normalize_text(notes),
                    now_text(),
                ),
            )
        except sqlite3.IntegrityError as exc:
            raise ERPError("Bu üretim emir numarası daha önce kullanılmış.") from exc

        order_id = cursor.lastrowid
        factor = planned_quantity / float(recipe["basis_quantity"])
        connection.executemany(
            """
            INSERT INTO production_order_materials
                (production_order_id, item_id, planned_quantity)
            VALUES (?, ?, ?)
            """,
            [(order_id, line["item_id"], float(line["quantity"]) * factor) for line in lines],
        )
    return order_no


def start_production_order(order_id):
    with db_session(write=True) as connection:
        cursor = connection.execute(
            """
            UPDATE production_orders
            SET status = 'IN_PROGRESS', started_at = ?
            WHERE id = ? AND status = 'PLANNED'
            """,
            (now_text(), order_id),
        )
        if cursor.rowcount != 1:
            raise ERPError("Yalnızca planlanmış bir üretim emri başlatılabilir.")


def plan_fifo_allocations(connection, item_id, warehouse_id, required_quantity):
    required_quantity = float(required_quantity)
    rows = connection.execute(
        """
        SELECT
            sm.lot_id,
            l.lot_no,
            SUM(sm.quantity_delta) AS balance,
            MIN(sm.movement_time) AS first_movement
        FROM stock_movements sm
        JOIN lots l ON l.id = sm.lot_id
        WHERE sm.item_id = ? AND sm.warehouse_id = ?
        GROUP BY sm.lot_id, l.lot_no
        HAVING SUM(sm.quantity_delta) > 0.000001
        ORDER BY first_movement, sm.lot_id
        """,
        (item_id, warehouse_id),
    ).fetchall()
    total_available = sum(float(row["balance"]) for row in rows)
    if total_available + 0.000001 < required_quantity:
        item = connection.execute("SELECT name, unit FROM items WHERE id = ?", (item_id,)).fetchone()
        raise ERPError(
            f"{item['name']} için stok yetersiz. Gereken: {required_quantity:,.3f} "
            f"{item['unit']}, mevcut: {total_available:,.3f} {item['unit']}."
        )
    remaining = required_quantity
    allocations = []
    for row in rows:
        if remaining <= 0.000001:
            break
        amount = min(float(row["balance"]), remaining)
        allocations.append((row["lot_id"], row["lot_no"], amount))
        remaining -= amount
    return allocations


def complete_production_order(
    order_id,
    actual_output_quantity,
    actual_consumptions,
    waste_quantity=0,
    waste_disposition="NONE",
    waste_type="",
    waste_reason="",
    waste_item_id=None,
    waste_lot_no="",
    waste_warehouse_id=None,
):
    gross_quantity = float(actual_output_quantity)
    waste_quantity = float(waste_quantity or 0)
    waste_disposition = normalize_code(waste_disposition) or "NONE"
    waste_type = normalize_text(waste_type)
    waste_reason = normalize_text(waste_reason)

    if gross_quantity <= 0:
        raise ERPError("Brüt üretim miktarı sıfırdan büyük olmalıdır.")
    if waste_quantity < 0:
        raise ERPError("Ziyan miktarı negatif olamaz.")
    if waste_quantity > gross_quantity:
        raise ERPError("Ziyan miktarı brüt üretim miktarından büyük olamaz.")

    net_output_quantity = gross_quantity - waste_quantity
    if waste_quantity <= 0.000001:
        waste_quantity = 0.0
        waste_disposition = "NONE"
        waste_type = ""
        waste_reason = ""
        waste_item_id = None
        waste_lot_no = ""
        waste_warehouse_id = None
    else:
        if waste_disposition not in {"RECOVERABLE", "DISPOSAL"}:
            raise ERPError("Ziyanın nasıl değerlendirileceği seçilmelidir.")
        if not waste_type:
            raise ERPError("Ziyan türü seçilmelidir.")
        if waste_disposition == "RECOVERABLE":
            waste_lot_no = normalize_code(waste_lot_no)
            if not waste_item_id or not waste_warehouse_id or not waste_lot_no:
                raise ERPError(
                    "Tekrar kullanılabilir ziyan için stok kartı, depo ve ziyan lotu zorunludur."
                )
        else:
            waste_item_id = None
            waste_lot_no = ""
            waste_warehouse_id = None

    with db_session(write=True) as connection:
        order = connection.execute(
            "SELECT * FROM production_orders WHERE id = ?", (order_id,)
        ).fetchone()
        if not order or order["status"] != "IN_PROGRESS":
            raise ERPError("Yalnızca üretimde durumundaki emir tamamlanabilir.")

        waste_item = None
        if waste_disposition == "RECOVERABLE":
            waste_item = connection.execute(
                "SELECT id, name, unit FROM items WHERE id = ? AND active = 1",
                (int(waste_item_id),),
            ).fetchone()
            waste_warehouse = connection.execute(
                "SELECT id FROM warehouses WHERE id = ? AND active = 1",
                (int(waste_warehouse_id),),
            ).fetchone()
            product = connection.execute(
                "SELECT unit FROM items WHERE id = ?", (order["product_item_id"],)
            ).fetchone()
            if not waste_item or not waste_warehouse:
                raise ERPError("Ziyan için seçilen stok kartı veya depo aktif değil.")
            if waste_item["unit"] != product["unit"]:
                raise ERPError(
                    "Ziyan stok kartının birimi ile üretilen ürünün birimi aynı olmalıdır."
                )
        requirements = connection.execute(
            """
            SELECT pom.item_id, pom.planned_quantity, i.name, i.unit
            FROM production_order_materials pom
            JOIN items i ON i.id = pom.item_id
            WHERE pom.production_order_id = ?
            ORDER BY i.category, i.name
            """,
            (order_id,),
        ).fetchall()

        all_allocations = {}
        for requirement in requirements:
            quantity = float(actual_consumptions.get(requirement["item_id"], 0))
            if quantity < 0:
                raise ERPError("Fiili tüketim miktarı negatif olamaz.")
            all_allocations[requirement["item_id"]] = plan_fifo_allocations(
                connection,
                requirement["item_id"],
                order["source_warehouse_id"],
                quantity,
            ) if quantity > 0 else []

        for requirement in requirements:
            item_id = requirement["item_id"]
            actual_quantity = float(actual_consumptions.get(item_id, 0))
            for lot_id, lot_no, allocation_quantity in all_allocations[item_id]:
                add_movement(
                    connection,
                    item_id=item_id,
                    lot_id=lot_id,
                    warehouse_id=order["source_warehouse_id"],
                    quantity_delta=-allocation_quantity,
                    movement_type="PRODUCTION_CONSUMPTION",
                    reference_type="PRODUCTION",
                    reference_no=order["order_no"],
                    notes=f"FIFO tüketim; kaynak lot: {lot_no}",
                )
            connection.execute(
                """
                UPDATE production_order_materials
                SET actual_quantity = ?
                WHERE production_order_id = ? AND item_id = ?
                """,
                (actual_quantity, order_id, item_id),
            )

        if net_output_quantity > 0.000001:
            product_lot_id = get_or_create_lot(
                connection, order["product_item_id"], order["product_lot_no"]
            )
            add_movement(
                connection,
                item_id=order["product_item_id"],
                lot_id=product_lot_id,
                warehouse_id=order["output_warehouse_id"],
                quantity_delta=net_output_quantity,
                movement_type="PRODUCTION_OUTPUT",
                reference_type="PRODUCTION",
                reference_no=order["order_no"],
                notes=(
                    f"Net satışa uygun mamul; brüt: {gross_quantity:.3f}, "
                    f"ziyan: {waste_quantity:.3f}"
                ),
            )

        if waste_disposition == "RECOVERABLE":
            waste_lot_id = get_or_create_lot(
                connection, int(waste_item_id), waste_lot_no
            )
            waste_note = f"{waste_type}; kaynak ürün lotu: {order['product_lot_no']}"
            if waste_reason:
                waste_note += f"; açıklama: {waste_reason}"
            add_movement(
                connection,
                item_id=int(waste_item_id),
                lot_id=waste_lot_id,
                warehouse_id=int(waste_warehouse_id),
                quantity_delta=waste_quantity,
                movement_type="PRODUCTION_WASTE_RECOVERY",
                reference_type="PRODUCTION",
                reference_no=order["order_no"],
                notes=waste_note,
            )

        connection.execute(
            """
            UPDATE production_orders
            SET status = 'COMPLETED', gross_quantity = ?, actual_quantity = ?,
                waste_quantity = ?, waste_disposition = ?, waste_type = ?,
                waste_reason = ?, waste_item_id = ?, waste_lot_no = ?,
                waste_warehouse_id = ?, completed_at = ?
            WHERE id = ?
            """,
            (
                gross_quantity,
                net_output_quantity,
                waste_quantity,
                waste_disposition,
                waste_type or None,
                waste_reason or None,
                int(waste_item_id) if waste_item_id else None,
                waste_lot_no or None,
                int(waste_warehouse_id) if waste_warehouse_id else None,
                now_text(),
                order_id,
            ),
        )


def reverse_reference_movements(connection, reference_type, reference_no, reversal_type, reason):
    originals = connection.execute(
        """
        SELECT * FROM stock_movements
        WHERE reference_type = ? AND reference_no = ? AND reversal_of IS NULL
        ORDER BY id
        """,
        (reference_type, reference_no),
    ).fetchall()
    if not originals:
        raise ERPError("Ters çevrilecek stok hareketi bulunamadı.")
    original_ids = [row["id"] for row in originals]
    placeholders = ",".join("?" for _ in original_ids)
    reversed_count = connection.execute(
        f"SELECT COUNT(*) AS count FROM stock_movements WHERE reversal_of IN ({placeholders})",
        tuple(original_ids),
    ).fetchone()["count"]
    if reversed_count:
        raise ERPError("Bu işlem daha önce iptal edilmiş.")

    for movement in originals:
        if float(movement["quantity_delta"]) > 0:
            current = get_balance(
                connection,
                movement["item_id"],
                movement["lot_id"],
                movement["warehouse_id"],
            )
            if current + 0.000001 < float(movement["quantity_delta"]):
                raise ERPError(
                    "İşlem çıktısının bir kısmı daha sonra kullanıldığı için iptal yapılamaz. "
                    "Önce bağlı sevkiyat veya tüketimler iptal edilmelidir."
                )

    for movement in originals:
        add_movement(
            connection,
            item_id=movement["item_id"],
            lot_id=movement["lot_id"],
            warehouse_id=movement["warehouse_id"],
            quantity_delta=-float(movement["quantity_delta"]),
            movement_type=reversal_type,
            reference_type=f"{reference_type}_CANCEL",
            reference_no=reference_no,
            notes=reason,
            reversal_of=movement["id"],
        )


def cancel_production_order(order_id, reason):
    reason = normalize_text(reason)
    if not reason:
        raise ERPError("İptal nedeni yazılmalıdır.")
    with db_session(write=True) as connection:
        order = connection.execute(
            "SELECT * FROM production_orders WHERE id = ?", (order_id,)
        ).fetchone()
        if not order or order["status"] == "CANCELLED":
            raise ERPError("Üretim emri bulunamadı veya zaten iptal edilmiş.")
        if order["status"] == "COMPLETED":
            reverse_reference_movements(
                connection,
                "PRODUCTION",
                order["order_no"],
                "PRODUCTION_CANCEL",
                reason,
            )
        connection.execute(
            """
            UPDATE production_orders
            SET status = 'CANCELLED', cancelled_at = ?,
                notes = CASE WHEN notes IS NULL OR notes = '' THEN ? ELSE notes || ' | İptal: ' || ? END
            WHERE id = ?
            """,
            (now_text(), reason, reason, order_id),
        )


def create_shipment(
    item_id,
    lot_id,
    warehouse_id,
    quantity,
    customer_name,
    dispatch_note_no,
    plate,
    notes,
):
    quantity = float(quantity)
    customer_name = normalize_text(customer_name)
    dispatch_note_no = normalize_code(dispatch_note_no)
    if quantity <= 0 or not customer_name or not dispatch_note_no:
        raise ERPError("Müşteri, irsaliye numarası ve sevk miktarı zorunludur.")
    shipment_no = make_reference("SVK")
    with db_session(write=True) as connection:
        item = connection.execute("SELECT category FROM items WHERE id = ?", (item_id,)).fetchone()
        if not item or item["category"] != "Mamul":
            raise ERPError("Yalnızca Mamul kategorisindeki ürünler sevk edilebilir.")
        balance = get_balance(connection, item_id, lot_id, warehouse_id)
        if quantity > balance + 0.000001:
            raise ERPError(f"Seçilen mamul lotunda yalnızca {balance:,.3f} stok var.")
        cursor = connection.execute(
            """
            INSERT INTO shipments
                (shipment_no, shipment_time, customer_name, dispatch_note_no,
                 plate, status, notes, created_at)
            VALUES (?, ?, ?, ?, ?, 'COMPLETED', ?, ?)
            """,
            (
                shipment_no,
                now_text(),
                customer_name,
                dispatch_note_no,
                normalize_code(plate),
                normalize_text(notes),
                now_text(),
            ),
        )
        shipment_id = cursor.lastrowid
        movement_id = add_movement(
            connection,
            item_id=item_id,
            lot_id=lot_id,
            warehouse_id=warehouse_id,
            quantity_delta=-quantity,
            movement_type="SHIPMENT",
            reference_type="SHIPMENT",
            reference_no=shipment_no,
            notes=f"{customer_name} | İrsaliye: {dispatch_note_no}",
        )
        connection.execute(
            """
            INSERT INTO shipment_lines
                (shipment_id, item_id, lot_id, warehouse_id, quantity, movement_id)
            VALUES (?, ?, ?, ?, ?, ?)
            """,
            (shipment_id, item_id, lot_id, warehouse_id, quantity, movement_id),
        )
    return shipment_no


def cancel_shipment(shipment_id, reason):
    reason = normalize_text(reason)
    if not reason:
        raise ERPError("Sevkiyat iptal nedeni yazılmalıdır.")
    with db_session(write=True) as connection:
        shipment = connection.execute(
            "SELECT * FROM shipments WHERE id = ?", (shipment_id,)
        ).fetchone()
        if not shipment or shipment["status"] != "COMPLETED":
            raise ERPError("Sevkiyat bulunamadı veya daha önce iptal edilmiş.")
        reverse_reference_movements(
            connection,
            "SHIPMENT",
            shipment["shipment_no"],
            "SHIPMENT_CANCEL",
            reason,
        )
        connection.execute(
            """
            UPDATE shipments
            SET status = 'CANCELLED', cancelled_at = ?,
                notes = CASE WHEN notes IS NULL OR notes = '' THEN ? ELSE notes || ' | İptal: ' || ? END
            WHERE id = ?
            """,
            (now_text(), reason, reason, shipment_id),
        )


def production_orders_df(status=None):
    params = []
    where_sql = ""
    if status:
        where_sql = "WHERE po.status = ?"
        params.append(status)
    df = query_df(
        f"""
        SELECT
            po.id,
            po.order_no AS "Emir No",
            po.planned_date AS "Plan Tarihi",
            po.recipe_name_snapshot AS "Reçete",
            po.recipe_version AS "Reçete Versiyonu",
            i.name AS "Ürün",
            po.product_lot_no AS "Ürün LOT",
            po.planned_quantity AS "Planlanan Miktar",
            po.gross_quantity AS "Brüt Üretim",
            po.actual_quantity AS "Net Mamul",
            po.waste_quantity AS "Ziyan",
            CASE
                WHEN COALESCE(po.gross_quantity, 0) > 0
                THEN ROUND(po.waste_quantity * 100.0 / po.gross_quantity, 4)
                ELSE 0
            END AS "Ziyan Oranı (%)",
            CASE po.waste_disposition
                WHEN 'RECOVERABLE' THEN 'Tekrar Kullanılabilir'
                WHEN 'DISPOSAL' THEN 'Kullanılamaz / Bertaraf'
                ELSE 'Ziyan Yok'
            END AS "Ziyan Durumu",
            po.waste_type AS "Ziyan Türü",
            po.waste_reason AS "Ziyan Açıklaması",
            wi.name AS "Ziyan Stok Kartı",
            po.waste_lot_no AS "Ziyan LOT",
            ww.name AS "Ziyan Deposu",
            i.unit AS "Birim",
            sw.name AS "Kaynak Depo",
            ow.name AS "Çıktı Deposu",
            po.status AS "Durum",
            po.notes AS "Açıklama"
        FROM production_orders po
        JOIN items i ON i.id = po.product_item_id
        JOIN warehouses sw ON sw.id = po.source_warehouse_id
        JOIN warehouses ow ON ow.id = po.output_warehouse_id
        LEFT JOIN items wi ON wi.id = po.waste_item_id
        LEFT JOIN warehouses ww ON ww.id = po.waste_warehouse_id
        {where_sql}
        ORDER BY po.id DESC
        """,
        tuple(params),
    )
    if not df.empty:
        df["Durum"] = df["Durum"].map(ORDER_STATUS_LABELS).fillna(df["Durum"])
    return df


def shipments_df(start_date=None, end_date=None):
    conditions = []
    params = []
    if start_date:
        conditions.append("date(s.shipment_time) >= date(?)")
        params.append(str(start_date))
    if end_date:
        conditions.append("date(s.shipment_time) <= date(?)")
        params.append(str(end_date))
    where_sql = f"WHERE {' AND '.join(conditions)}" if conditions else ""
    df = query_df(
        f"""
        SELECT
            s.id,
            s.shipment_no AS "Sevkiyat No",
            s.shipment_time AS "Tarih",
            s.customer_name AS "Müşteri",
            s.dispatch_note_no AS "İrsaliye No",
            s.plate AS "Plaka",
            i.name AS "Ürün",
            l.lot_no AS "LOT No",
            w.name AS "Depo",
            sl.quantity AS "Miktar",
            i.unit AS "Birim",
            s.status AS "Durum",
            s.notes AS "Açıklama"
        FROM shipments s
        JOIN shipment_lines sl ON sl.shipment_id = s.id
        JOIN items i ON i.id = sl.item_id
        JOIN lots l ON l.id = sl.lot_id
        JOIN warehouses w ON w.id = sl.warehouse_id
        {where_sql}
        ORDER BY s.id DESC
        """,
        tuple(params),
    )
    if not df.empty:
        df["Durum"] = df["Durum"].map(
            {"COMPLETED": "Tamamlandı", "CANCELLED": "İptal Edildi"}
        ).fillna(df["Durum"])
    return df


def recipe_summary_df():
    return query_df(
        """
        SELECT
            r.id,
            r.name AS "Reçete Adı",
            r.recipe_type AS "Sınıf",
            p.name AS "Üretilecek Ürün",
            r.basis_quantity AS "Baz Miktar",
            p.unit AS "Ürün Birimi",
            r.version AS "Versiyon",
            CASE WHEN r.active = 1 THEN 'Aktif' ELSE 'Pasif' END AS "Durum",
            GROUP_CONCAT(i.name || ': ' || ROUND(rl.quantity, 6) || ' ' || i.unit, ' | ') AS "Tüketim Kalemleri"
        FROM recipes r
        JOIN items p ON p.id = r.product_item_id
        LEFT JOIN recipe_lines rl ON rl.recipe_id = r.id
        LEFT JOIN items i ON i.id = rl.item_id
        GROUP BY r.id, r.name, r.recipe_type, p.name, r.basis_quantity,
                 p.unit, r.version, r.active
        ORDER BY r.name
        """
    )


def build_excel_report(start_date, end_date):
    stock = stock_balance_df(positive_only=False)
    movements = movement_history_df(start_date, end_date)
    orders = production_orders_df()
    shipments = shipments_df(start_date, end_date)
    recipes = recipe_summary_df()

    for frame in [stock, movements, orders, shipments, recipes]:
        for hidden in ["item_id", "lot_id", "warehouse_id", "id"]:
            if hidden in frame.columns:
                frame.drop(columns=[hidden], inplace=True)

    buffer = io.BytesIO()
    with pd.ExcelWriter(buffer, engine="openpyxl") as writer:
        stock.to_excel(writer, index=False, sheet_name="Anlık Stok")
        movements.to_excel(writer, index=False, sheet_name="Dönem Hareketleri")
        orders.to_excel(writer, index=False, sheet_name="Üretim Emirleri")
        shipments.to_excel(writer, index=False, sheet_name="Sevkiyatlar")
        recipes.to_excel(writer, index=False, sheet_name="Reçeteler")

        from openpyxl.styles import Alignment, Font, PatternFill

        header_fill = PatternFill("solid", fgColor="17365D")
        header_font = Font(color="FFFFFF", bold=True)
        for worksheet in writer.book.worksheets:
            worksheet.freeze_panes = "A2"
            worksheet.auto_filter.ref = worksheet.dimensions
            for cell in worksheet[1]:
                cell.fill = header_fill
                cell.font = header_font
                cell.alignment = Alignment(horizontal="center", vertical="center")
            for column_cells in worksheet.columns:
                max_length = max(
                    len(str(cell.value)) if cell.value is not None else 0
                    for cell in column_cells
                )
                worksheet.column_dimensions[column_cells[0].column_letter].width = min(
                    max(max_length + 2, 12), 45
                )
            worksheet.sheet_view.showGridLines = False
    return buffer.getvalue()


def build_database_backup():
    if not DB_PATH.exists():
        raise ERPError("Yedeklenecek veritabanı bulunamadı.")
    with tempfile.NamedTemporaryFile(suffix=".db") as temporary_file:
        source = sqlite3.connect(str(DB_PATH))
        destination = sqlite3.connect(temporary_file.name)
        try:
            source.backup(destination)
            destination.commit()
        finally:
            destination.close()
            source.close()
        temporary_file.seek(0)
        return temporary_file.read()


def active_items():
    return query_rows(
        """
        SELECT id, code, name, category, unit, min_stock, active
        FROM items WHERE active = 1
        ORDER BY CASE category
            WHEN 'Hammadde' THEN 1
            WHEN 'Yardımcı Kimyasal' THEN 2
            WHEN 'Ambalaj' THEN 3
            WHEN 'Ara Mamul' THEN 4
            WHEN 'Mamul' THEN 5
            ELSE 6 END, name
        """
    )


def active_warehouses():
    return query_rows(
        "SELECT id, code, name, warehouse_type FROM warehouses WHERE active = 1 ORDER BY name"
    )


def option_map(rows, label_builder):
    return {row["id"]: label_builder(row) for row in rows}


def display_frame(df, hidden_columns=None):
    frame = df.copy()
    for column in hidden_columns or []:
        if column in frame.columns:
            frame.drop(columns=[column], inplace=True)
    st.dataframe(frame, use_container_width=True, hide_index=True)


def show_error(exc):
    if isinstance(exc, ERPError):
        st.error(f"❌ {exc}")
    else:
        st.error(f"❌ İşlem tamamlanamadı: {exc}")


def dashboard_page():
    st.header("📊 Fabrika Stok ve Üretim Kokpiti")
    warehouses = active_warehouses()
    warehouse_options = {0: "Tüm Depolar"}
    warehouse_options.update(
        {row["id"]: f"{row['code']} | {row['name']}" for row in warehouses}
    )
    selected_warehouse_id = st.selectbox(
        "Görüntülenecek Depo",
        list(warehouse_options),
        format_func=lambda value: warehouse_options[value],
        key="dashboard_warehouse_filter",
    )
    warehouse_id = selected_warehouse_id or None
    selected_warehouse_name = warehouse_options[selected_warehouse_id]

    today = str(date.today())
    stock_condition = ""
    shipment_condition = ""
    metric_params = []
    if warehouse_id:
        stock_condition = "AND sm.warehouse_id = ?"
        metric_params.append(warehouse_id)
    metric_params.append(today)
    if warehouse_id:
        shipment_condition = "AND sl.warehouse_id = ?"
        metric_params.append(warehouse_id)
    metrics = query_rows(
        f"""
        SELECT
            (SELECT COUNT(*) FROM items WHERE active = 1) AS item_count,
            (SELECT COUNT(*) FROM production_orders WHERE status IN ('PLANNED', 'IN_PROGRESS')) AS open_orders,
            (SELECT COALESCE(SUM(sm.quantity_delta), 0)
             FROM stock_movements sm JOIN items i ON i.id = sm.item_id
             WHERE i.category = 'Mamul' AND i.unit = 'Kg' {stock_condition}) AS finished_kg,
            (SELECT COALESCE(SUM(sl.quantity), 0)
             FROM shipments s
             JOIN shipment_lines sl ON sl.shipment_id = s.id
             JOIN items i ON i.id = sl.item_id
             WHERE date(s.shipment_time) = date(?)
               AND s.status = 'COMPLETED' AND i.unit = 'Kg' {shipment_condition}) AS shipped_today_kg
        """,
        tuple(metric_params),
    )[0]
    c1, c2, c3, c4 = st.columns(4)
    c1.metric("Aktif Stok Kartı", f"{int(metrics['item_count'])}")
    c2.metric("Açık Üretim Emri", f"{int(metrics['open_orders'])}")
    c3.metric(f"Satışa Hazır Mamul · {selected_warehouse_name}", f"{float(metrics['finished_kg']):,.1f} Kg")
    c4.metric(f"Bugünkü Sevkiyat · {selected_warehouse_name}", f"{float(metrics['shipped_today_kg']):,.1f} Kg")

    st.subheader(f"Kritik Stok Uyarıları · {selected_warehouse_name}")
    movement_join = "LEFT JOIN stock_movements sm ON sm.item_id = i.id"
    low_stock_params = ()
    if warehouse_id:
        movement_join += " AND sm.warehouse_id = ?"
        low_stock_params = (warehouse_id,)
    low_stock = query_df(
        f"""
        SELECT
            i.code AS "Stok Kodu",
            i.name AS "Malzeme / Ürün",
            i.category AS "Kategori",
            ROUND(COALESCE(SUM(sm.quantity_delta), 0), 6) AS "Toplam Stok",
            i.min_stock AS "Minimum Stok",
            i.unit AS "Birim"
        FROM items i
        {movement_join}
        WHERE i.active = 1 AND i.min_stock > 0
        GROUP BY i.id, i.code, i.name, i.category, i.min_stock, i.unit
        HAVING COALESCE(SUM(sm.quantity_delta), 0) < i.min_stock
        ORDER BY (i.min_stock - COALESCE(SUM(sm.quantity_delta), 0)) DESC
        """,
        low_stock_params,
    )
    if low_stock.empty:
        st.success("✅ Minimum stok seviyesinin altında kalan aktif kart bulunmuyor.")
    else:
        st.warning(f"⚠️ {len(low_stock)} stok kartı minimum seviyenin altında.")
        display_frame(low_stock)

    stock_category_options = {
        "Tümü": "Tüm Kategoriler",
        "Hammadde": "Hammadde",
        "Yardımcı Kimyasal": "Yardımcı Kimyasallar",
        "Ambalaj": "Ambalaj",
        "Ara Mamul": "Ara Mamul",
        "Mamul": "Mamul",
    }
    selected_stock_category = st.selectbox(
        "Anlık Stok Kategorisi",
        list(stock_category_options),
        format_func=lambda value: stock_category_options[value],
        key="dashboard_stock_category_filter",
    )
    selected_categories = (
        None if selected_stock_category == "Tümü" else [selected_stock_category]
    )
    selected_category_name = stock_category_options[selected_stock_category]

    st.subheader(
        f"Anlık Lot Bazlı Stok · {selected_warehouse_name} · {selected_category_name}"
    )
    stock = stock_balance_df(
        positive_only=True,
        categories=selected_categories,
        warehouse_id=warehouse_id,
    )
    if stock.empty:
        st.info("Henüz stok hareketi bulunmuyor. İlk kaydı Stok Yönetimi > Mal Kabul bölümünden yapabilirsiniz.")
    else:
        display_frame(stock, ["item_id", "lot_id", "warehouse_id"])


def stock_page():
    st.header("📦 Stok Yönetimi")
    tab_receipt, tab_transfer, tab_adjustment, tab_stock, tab_movements = st.tabs(
        ["📥 Mal Kabul", "🔄 Transfer", "📉 Fire / Düzeltme", "📦 Stoklar", "📋 Hareketler"]
    )

    items = active_items()
    warehouses = active_warehouses()
    item_labels = option_map(items, lambda row: f"{row['code']} | {row['name']} ({row['unit']})")
    warehouse_labels = option_map(warehouses, lambda row: f"{row['code']} | {row['name']}")

    with tab_receipt:
        if not items or not warehouses:
            st.warning("Mal kabul için en az bir aktif stok kartı ve depo gereklidir.")
        else:
            selected_item_id = st.selectbox(
                "Malzeme / Stok Kartı",
                list(item_labels),
                format_func=lambda value: item_labels[value],
                key="receipt_item",
            )
            selected_item = next(row for row in items if row["id"] == selected_item_id)
            with st.form("receipt_form", clear_on_submit=True):
                c1, c2 = st.columns(2)
                movement_date = c1.date_input("Giriş Tarihi", value=date.today())
                warehouse_id = c2.selectbox(
                    "Giriş Deposu",
                    list(warehouse_labels),
                    format_func=lambda value: warehouse_labels[value],
                )
                lot_no = st.text_input("Gelen LOT / Parti Numarası")
                quantity = st.number_input(
                    f"Gelen Miktar ({selected_item['unit']})",
                    min_value=0.000001,
                    value=1000.0 if selected_item["unit"] == "Kg" else 1.0,
                    step=1.0,
                    format="%.3f",
                )
                document_no = st.text_input("İrsaliye / Giriş Belgesi No (boşsa otomatik oluşur)")
                notes = st.text_area("Açıklama")
                submitted = st.form_submit_button("Malzemeyi Depoya Kabul Et", use_container_width=True)
            if submitted:
                try:
                    timestamp = datetime.combine(movement_date, datetime.now().time()).strftime("%Y-%m-%d %H:%M:%S")
                    ref = create_receipt(
                        selected_item_id,
                        warehouse_id,
                        lot_no,
                        quantity,
                        timestamp,
                        document_no,
                        notes,
                    )
                    st.success(f"✅ Mal kabul kaydedildi. Referans: {ref}")
                    st.rerun()
                except Exception as exc:
                    show_error(exc)

    with tab_transfer:
        available = stock_balance_df(positive_only=True)
        if available.empty or not warehouses:
            st.info("Transfer edilebilecek pozitif stok bulunmuyor.")
        else:
            row_options = {}
            for _, row in available.iterrows():
                label = (
                    f"{row['Stok Kodu']} | {row['Malzeme / Ürün']} | LOT {row['LOT No']} | "
                    f"{row['Depo']} | {row['Güncel Stok']:,.3f} {row['Birim']}"
                )
                row_options[label] = row
            selected_label = st.selectbox("Kaynak Stok", list(row_options), key="transfer_source")
            selected = row_options[selected_label]
            targets = [row for row in warehouses if row["id"] != int(selected["warehouse_id"])]
            target_labels = option_map(targets, lambda row: f"{row['code']} | {row['name']}")
            if not target_labels:
                st.warning("Transfer için kaynak depodan farklı en az bir aktif hedef depo gereklidir.")
            else:
                with st.form("transfer_form", clear_on_submit=True):
                    target_id = st.selectbox(
                        "Hedef Depo",
                        list(target_labels),
                        format_func=lambda value: target_labels[value],
                    )
                    quantity = st.number_input(
                        f"Transfer Miktarı ({selected['Birim']})",
                        min_value=0.000001,
                        max_value=float(selected["Güncel Stok"]),
                        value=float(min(50.0, selected["Güncel Stok"])),
                        step=1.0,
                        format="%.3f",
                    )
                    notes = st.text_area("Transfer Açıklaması")
                    submitted = st.form_submit_button("Transferi Onayla", use_container_width=True)
                if submitted:
                    try:
                        ref = transfer_stock(
                            int(selected["item_id"]),
                            int(selected["lot_id"]),
                            int(selected["warehouse_id"]),
                            target_id,
                            quantity,
                            notes,
                        )
                        st.success(f"✅ Transfer tamamlandı. Referans: {ref}")
                        st.rerun()
                    except Exception as exc:
                        show_error(exc)

    with tab_adjustment:
        adjustment_type = st.radio(
            "İşlem Türü",
            ["Stok Düşümü / Fire", "Stok İlavesi / Sayım Fazlası"],
            horizontal=True,
        )
        if adjustment_type == "Stok Düşümü / Fire":
            available = stock_balance_df(positive_only=True)
            if available.empty:
                st.info("Düşülebilecek pozitif stok bulunmuyor.")
            else:
                row_options = {}
                for _, row in available.iterrows():
                    label = (
                        f"{row['Stok Kodu']} | {row['Malzeme / Ürün']} | LOT {row['LOT No']} | "
                        f"{row['Depo']} | {row['Güncel Stok']:,.3f} {row['Birim']}"
                    )
                    row_options[label] = row
                selected_label = st.selectbox("Düşülecek Stok", list(row_options), key="adjust_out_stock")
                selected = row_options[selected_label]
                with st.form("adjust_out_form", clear_on_submit=True):
                    quantity = st.number_input(
                        f"Düşülecek Miktar ({selected['Birim']})",
                        min_value=0.000001,
                        max_value=float(selected["Güncel Stok"]),
                        value=float(min(10.0, selected["Güncel Stok"])),
                        step=1.0,
                        format="%.3f",
                    )
                    reason = st.selectbox(
                        "Düşüm Nedeni",
                        [
                            "Kullanım Süresi Dolması",
                            "Saha Firesi / Dökülme",
                            "Laboratuvar Analiz Sarfiyatı",
                            "Sayım Eksiği",
                            "Düzeltme Fişi",
                        ],
                    )
                    submitted = st.form_submit_button("Stok Düşümünü Onayla", use_container_width=True)
                if submitted:
                    try:
                        ref = create_adjustment(
                            item_id=int(selected["item_id"]),
                            warehouse_id=int(selected["warehouse_id"]),
                            lot_id=int(selected["lot_id"]),
                            quantity=quantity,
                            direction="OUT",
                            reason=reason,
                        )
                        st.success(f"✅ Stok düşümü kaydedildi. Referans: {ref}")
                        st.rerun()
                    except Exception as exc:
                        show_error(exc)
        else:
            if not items or not warehouses:
                st.warning("İşlem için aktif stok kartı ve depo gereklidir.")
            else:
                with st.form("adjust_in_form", clear_on_submit=True):
                    item_id = st.selectbox(
                        "Malzeme / Stok Kartı",
                        list(item_labels),
                        format_func=lambda value: item_labels[value],
                    )
                    warehouse_id = st.selectbox(
                        "Depo",
                        list(warehouse_labels),
                        format_func=lambda value: warehouse_labels[value],
                    )
                    lot_no = st.text_input("LOT / Parti Numarası")
                    quantity = st.number_input("İlave Miktarı", min_value=0.001, value=1.0, step=1.0, format="%.3f")
                    reason = st.text_input("Düzeltme Nedeni", value="Sayım Fazlası")
                    submitted = st.form_submit_button("Stok İlavesini Onayla", use_container_width=True)
                if submitted:
                    try:
                        ref = create_adjustment(
                            item_id=item_id,
                            warehouse_id=warehouse_id,
                            lot_no=lot_no,
                            quantity=quantity,
                            direction="IN",
                            reason=reason,
                        )
                        st.success(f"✅ Stok ilavesi kaydedildi. Referans: {ref}")
                        st.rerun()
                    except Exception as exc:
                        show_error(exc)

    with tab_stock:
        category = st.selectbox("Kategori Filtresi", ["Tümü"] + CATEGORIES, key="stock_category_filter")
        categories = None if category == "Tümü" else [category]
        stock = stock_balance_df(positive_only=False, categories=categories)
        if stock.empty:
            st.info("Seçilen filtreye uygun stok bulunmuyor.")
        else:
            display_frame(stock, ["item_id", "lot_id", "warehouse_id"])

    with tab_movements:
        c1, c2 = st.columns(2)
        start_date = c1.date_input("Başlangıç", value=date(date.today().year, 1, 1), key="mov_start")
        end_date = c2.date_input("Bitiş", value=date.today(), key="mov_end")
        movement_item_options = {0: "Tüm Stok Kartları", **item_labels}
        selected_item = st.selectbox(
            "Stok Kartı Filtresi",
            list(movement_item_options),
            format_func=lambda value: movement_item_options[value],
            key="movement_item_filter",
        )
        history = movement_history_df(start_date, end_date, selected_item or None)
        if history.empty:
            st.info("Seçilen aralıkta hareket bulunmuyor.")
        else:
            display_frame(history)


def recipe_editor():
    st.info(
        "Reçete miktarları, seçilen baz ürün miktarı için tanımlanır. "
        "Örneğin 1.000 Kg ürün bazı için PTA 850 Kg girilebilir."
    )
    tab_intermediate, tab_finished = st.tabs(
        ["⚙️ Ara Mamul Reçeteleri", "📦 Mamul Reçeteleri"]
    )
    with tab_intermediate:
        recipe_editor_for_type("Ara Mamul Reçetesi", "ara_mamul")
    with tab_finished:
        recipe_editor_for_type("Mamul Reçetesi", "mamul")


def recipe_editor_for_type(recipe_type, key_prefix):
    recipes = recipe_summary_df()
    if not recipes.empty:
        recipes = recipes[recipes["Sınıf"] == recipe_type].copy()
    st.subheader(recipe_type.replace(" Reçetesi", " Reçeteleri"))
    mode = st.radio(
        "İşlem",
        ["Yeni Reçete", "Mevcut Reçeteyi Düzenle"],
        horizontal=True,
        key=f"recipe_mode_{key_prefix}",
    )
    recipe_id = None
    existing = None
    existing_lines = {}
    if mode == "Mevcut Reçeteyi Düzenle":
        if recipes.empty:
            st.info("Düzenlenecek reçete bulunmuyor.")
            return
        recipe_options = dict(zip(recipes["id"].astype(int), recipes["Reçete Adı"]))
        recipe_id = st.selectbox(
            "Reçete",
            list(recipe_options),
            format_func=lambda value: recipe_options[value],
            key=f"recipe_edit_select_{key_prefix}",
        )
        existing = query_rows(
            """
            SELECT r.*, i.name AS product_name, i.unit AS product_unit
            FROM recipes r
            JOIN items i ON i.id = r.product_item_id
            WHERE r.id = ?
            """,
            (recipe_id,),
        )[0]
        existing_lines = {
            row["item_id"]: row["quantity"]
            for row in query_rows("SELECT item_id, quantity FROM recipe_lines WHERE recipe_id = ?", (recipe_id,))
        }

    items = active_items()
    known_item_ids = {row["id"] for row in items}
    missing_item_ids = [item_id for item_id in existing_lines if item_id not in known_item_ids]
    if missing_item_ids:
        placeholders = ",".join("?" for _ in missing_item_ids)
        items.extend(
            query_rows(
                f"""
                SELECT id, code, name, category, unit, min_stock, active
                FROM items WHERE id IN ({placeholders}) ORDER BY name
                """,
                tuple(missing_item_ids),
            )
        )
    if not items:
        st.warning("Önce stok kartı tanımlanmalıdır.")
        return
    item_labels = option_map(items, lambda row: f"{row['code']} | {row['name']} ({row['unit']})")

    default_materials = list(existing_lines) if existing else []
    selected_materials = st.multiselect(
        "Reçetede Tüketilecek Kalemler",
        list(item_labels),
        default=default_materials,
        format_func=lambda value: item_labels[value],
        key=f"recipe_materials_{key_prefix}_{recipe_id or 'new'}",
    )

    with st.form(f"recipe_form_{key_prefix}_{recipe_id or 'new'}"):
        name = st.text_input(
            "Reçete Adı",
            value=existing["name"] if existing else "",
            key=f"recipe_name_{key_prefix}_{recipe_id or 'new'}",
        )
        target_category = "Ara Mamul" if recipe_type == "Ara Mamul Reçetesi" else "Mamul"
        st.caption(f"Reçete sınıfı: {recipe_type} · Ürün stok kategorisi: {target_category}")
        st.markdown("**Üretilecek ürün bilgisi**")
        c_product_name, c_product_unit = st.columns([3, 1])
        product_name = c_product_name.text_input(
            "Üretilecek Ürün Adı (serbest yazabilirsiniz)",
            value=existing["product_name"] if existing else "",
            placeholder="Örnek: Standart PET Resin 0.80 IV",
            key=f"recipe_product_name_{key_prefix}_{recipe_id or 'new'}",
        )
        default_product_unit = existing["product_unit"] if existing else "Kg"
        product_unit = c_product_unit.selectbox(
            "Ürün Birimi",
            UNITS,
            index=UNITS.index(default_product_unit),
            key=f"recipe_product_unit_{key_prefix}_{recipe_id or 'new'}",
        )
        st.caption(
            "Bu ürün sistemde yoksa reçete kaydedilirken stok kartı otomatik oluşturulur. "
            "Reçete sınıfına göre Ara Mamul veya Mamul kategorisine bağlanır."
        )
        basis_quantity = st.number_input(
            "Baz Üretim Miktarı",
            min_value=0.001,
            value=float(existing["basis_quantity"]) if existing else 1000.0,
            step=1.0,
            format="%.3f",
            help="Miktarın birimi yukarıda seçtiğiniz Ürün Birimidir.",
            key=f"recipe_basis_{key_prefix}_{recipe_id or 'new'}",
        )
        material_quantities = {}
        if selected_materials:
            st.markdown("**Baz üretim için tüketimler**")
        for item_id in selected_materials:
            item = next(row for row in items if row["id"] == item_id)
            material_quantities[item_id] = st.number_input(
                f"{item['name']} ({item['unit']})",
                min_value=0.0,
                value=float(existing_lines.get(item_id, 0.0)),
                step=0.001,
                format="%.6f",
                key=f"recipe_qty_{key_prefix}_{recipe_id or 'new'}_{item_id}",
            )
        submitted = st.form_submit_button("Reçeteyi Kaydet", use_container_width=True)
    if submitted:
        try:
            result = save_recipe(
                recipe_id,
                name,
                recipe_type,
                product_name,
                product_unit,
                basis_quantity,
                material_quantities,
            )
            if result["product_created"]:
                st.success(
                    f"✅ Reçete kaydedildi. '{result['product_name']}' ürünü için "
                    f"{result['product_code']} kodlu stok kartı otomatik oluşturuldu."
                )
            else:
                st.success(
                    "✅ Reçete kaydedildi ve mevcut ürün stok kartına bağlandı. "
                    "Açılmış üretim emirlerinin reçete kopyası değişmedi."
                )
            st.rerun()
        except Exception as exc:
            show_error(exc)

    st.subheader(f"Kayıtlı {recipe_type.replace(' Reçetesi', ' Reçeteleri')}")
    display_frame(recipes, ["id"])
    st.caption(
        "Varsayılan Standart Amorf Chips reçetesi, önceki kodundaki örnek katsayıların "
        "1.000 Kg baza çevrilmiş halidir; gerçek üretimde kullanmadan önce saha reçetesiyle doğrulayın."
    )


def production_page():
    st.header("🏭 Üretim Yönetimi")
    tab_new, tab_tracking, tab_complete, tab_recipes = st.tabs(
        ["➕ Yeni Emir", "📋 Emir Takibi", "✅ Üretimi Tamamla", "📝 Reçeteler"]
    )

    with tab_new:
        recipes = query_rows(
            """
            SELECT r.id, r.name, r.recipe_type, r.basis_quantity,
                   i.name AS product_name, i.unit AS product_unit
            FROM recipes r JOIN items i ON i.id = r.product_item_id
            WHERE r.active = 1 AND i.active = 1
            ORDER BY r.name
            """
        )
        warehouses = active_warehouses()
        if not recipes or not warehouses:
            st.warning("Üretim emri için aktif reçete ve depo gereklidir.")
        else:
            recipe_labels = option_map(
                recipes,
                lambda row: f"{row['recipe_type']} | {row['name']} → {row['product_name']}",
            )
            warehouse_labels = option_map(warehouses, lambda row: f"{row['code']} | {row['name']}")
            selected_recipe_id = st.selectbox(
                "Üretim Reçetesi",
                list(recipe_labels),
                format_func=lambda value: recipe_labels[value],
                key="new_order_recipe",
            )
            selected_recipe = next(row for row in recipes if row["id"] == selected_recipe_id)
            with st.form("new_order_form", clear_on_submit=True):
                c1, c2 = st.columns(2)
                order_no = c1.text_input("Üretim Emir No", value=make_reference("UE"))
                planned_date = c2.date_input("Planlanan Üretim Tarihi", value=date.today())
                product_lot_no = st.text_input("Üretilecek Ürün LOT No", value=f"LOT-{datetime.now().strftime('%Y%m%d%H%M')}")
                planned_quantity = st.number_input(
                    f"Planlanan Üretim Miktarı ({selected_recipe['product_unit']})",
                    min_value=0.000001,
                    value=1000.0,
                    step=1.0,
                    format="%.3f",
                )
                c3, c4 = st.columns(2)
                source_warehouse_id = c3.selectbox(
                    "Tüketim Kaynak Deposu",
                    list(warehouse_labels),
                    format_func=lambda value: warehouse_labels[value],
                )
                output_warehouse_id = c4.selectbox(
                    "Üretim Çıktı Deposu / Silo",
                    list(warehouse_labels),
                    format_func=lambda value: warehouse_labels[value],
                )
                notes = st.text_area("Üretim Notu")
                submitted = st.form_submit_button("Üretim Emrini Planla", use_container_width=True)
            if submitted:
                try:
                    ref = create_production_order(
                        order_no,
                        selected_recipe_id,
                        planned_quantity,
                        product_lot_no,
                        source_warehouse_id,
                        output_warehouse_id,
                        planned_date,
                        notes,
                    )
                    st.success(f"✅ {ref} numaralı üretim emri planlandı. Henüz stok düşülmedi.")
                    st.rerun()
                except Exception as exc:
                    show_error(exc)

    with tab_tracking:
        orders_raw = query_rows(
            """
            SELECT po.id, po.order_no, po.status, po.recipe_name_snapshot,
                   po.planned_quantity, po.actual_quantity, i.unit,
                   po.product_lot_no, po.planned_date
            FROM production_orders po JOIN items i ON i.id = po.product_item_id
            ORDER BY po.id DESC
            """
        )
        if not orders_raw:
            st.info("Kayıtlı üretim emri bulunmuyor.")
        else:
            status_filter = st.selectbox(
                "Durum Filtresi",
                ["Tümü", "PLANNED", "IN_PROGRESS", "COMPLETED", "CANCELLED"],
                format_func=lambda value: "Tümü" if value == "Tümü" else ORDER_STATUS_LABELS[value],
            )
            orders = production_orders_df(None if status_filter == "Tümü" else status_filter)
            display_frame(orders, ["id"])

            actionable = [row for row in orders_raw if row["status"] != "CANCELLED"]
            if actionable:
                order_labels = {
                    row["id"]: f"{row['order_no']} | {row['recipe_name_snapshot']} | {ORDER_STATUS_LABELS[row['status']]}"
                    for row in actionable
                }
                selected_order_id = st.selectbox(
                    "İşlem Yapılacak Emir",
                    list(order_labels),
                    format_func=lambda value: order_labels[value],
                    key="tracking_order_select",
                )
                selected_order = next(row for row in actionable if row["id"] == selected_order_id)
                c1, c2 = st.columns(2)
                if selected_order["status"] == "PLANNED":
                    if c1.button("▶️ Üretimi Başlat", use_container_width=True):
                        try:
                            start_production_order(selected_order_id)
                            st.success("✅ Üretim emri Üretimde durumuna alındı.")
                            st.rerun()
                        except Exception as exc:
                            show_error(exc)
                elif selected_order["status"] == "IN_PROGRESS":
                    c1.info("Fiili girişler için Üretimi Tamamla sekmesini kullanın.")
                else:
                    c1.info("Tamamlanan emir yalnızca bağlı stok hâlâ mevcutsa iptal edilebilir.")

                if c2.button("⚠️ İptal İşlemini Aç", use_container_width=True):
                    st.session_state["confirm_order_cancel"] = selected_order_id
                if st.session_state.get("confirm_order_cancel") == selected_order_id:
                    st.warning("Bu işlem kayıtları silmez; tamamlanan emirde ters stok hareketleri oluşturur.")
                    cancel_reason = st.text_input("Üretim Emri İptal Nedeni", key=f"order_cancel_reason_{selected_order_id}")
                    cc1, cc2 = st.columns(2)
                    if cc1.button("İptali Kesinleştir", key=f"confirm_cancel_order_{selected_order_id}", use_container_width=True):
                        try:
                            cancel_production_order(selected_order_id, cancel_reason)
                            st.session_state.pop("confirm_order_cancel", None)
                            st.success("✅ Üretim emri güvenli biçimde iptal edildi.")
                            st.rerun()
                        except Exception as exc:
                            show_error(exc)
                    if cc2.button("Vazgeç", key=f"abort_cancel_order_{selected_order_id}", use_container_width=True):
                        st.session_state.pop("confirm_order_cancel", None)
                        st.rerun()

    with tab_complete:
        in_progress = query_rows(
            """
            SELECT po.*, i.name AS product_name, i.unit AS product_unit,
                   w.name AS source_warehouse
            FROM production_orders po
            JOIN items i ON i.id = po.product_item_id
            JOIN warehouses w ON w.id = po.source_warehouse_id
            WHERE po.status = 'IN_PROGRESS'
            ORDER BY po.id DESC
            """
        )
        if not in_progress:
            st.info("Tamamlanmayı bekleyen Üretimde durumunda emir bulunmuyor.")
        else:
            order_labels = {
                row["id"]: f"{row['order_no']} | {row['recipe_name_snapshot']} | {row['product_lot_no']}"
                for row in in_progress
            }
            selected_order_id = st.selectbox(
                "Tamamlanacak Emir",
                list(order_labels),
                format_func=lambda value: order_labels[value],
                key="complete_order_select",
            )
            order = next(row for row in in_progress if row["id"] == selected_order_id)
            requirements = query_rows(
                """
                SELECT pom.item_id, pom.planned_quantity, i.name, i.unit,
                       COALESCE((
                           SELECT SUM(sm.quantity_delta)
                           FROM stock_movements sm
                           WHERE sm.item_id = pom.item_id
                             AND sm.warehouse_id = po.source_warehouse_id
                       ), 0) AS available
                FROM production_order_materials pom
                JOIN production_orders po ON po.id = pom.production_order_id
                JOIN items i ON i.id = pom.item_id
                WHERE pom.production_order_id = ?
                ORDER BY i.category, i.name
                """,
                (selected_order_id,),
            )
            st.caption(
                f"Kaynak depo: {order['source_warehouse']} · Ürün LOT: {order['product_lot_no']} · "
                "Tüketilecek hammaddeler mevcut lotlardan FIFO sırasıyla seçilir."
            )
            preview = pd.DataFrame(
                [
                    {
                        "Malzeme": row["name"],
                        "Planlanan": row["planned_quantity"],
                        "Mevcut Stok": row["available"],
                        "Birim": row["unit"],
                    }
                    for row in requirements
                ]
            )
            display_frame(preview)
            st.markdown("**Üretim Çıktısı ve Ziyan**")
            output_c1, output_c2 = st.columns(2)
            gross_output = output_c1.number_input(
                f"Brüt Gerçekleşen Üretim ({order['product_unit']})",
                min_value=0.001,
                value=float(order["planned_quantity"]),
                step=1.0,
                format="%.3f",
                key=f"gross_output_{selected_order_id}",
                help="Satışa uygun mamul ile oluşan ziyanın toplamıdır.",
            )
            waste_quantity = output_c2.number_input(
                f"Ziyan Miktarı ({order['product_unit']})",
                min_value=0.0,
                value=0.0,
                step=1.0,
                format="%.3f",
                key=f"waste_quantity_{selected_order_id}",
            )
            net_output = float(gross_output) - float(waste_quantity)
            waste_rate = (
                float(waste_quantity) * 100.0 / float(gross_output)
                if float(gross_output) > 0
                else 0.0
            )
            metric_c1, metric_c2, metric_c3, metric_c4 = st.columns(4)
            metric_c1.metric("Brüt Üretim", f"{float(gross_output):,.3f} {order['product_unit']}")
            metric_c2.metric("Ziyan", f"{float(waste_quantity):,.3f} {order['product_unit']}")
            metric_c3.metric("Net Mamul", f"{max(net_output, 0):,.3f} {order['product_unit']}")
            metric_c4.metric("Ziyan Oranı", f"%{waste_rate:,.3f}")
            st.caption("Kontrol: Brüt Üretim = Net Mamul + Ziyan")
            if net_output < 0:
                st.error("Ziyan miktarı brüt üretim miktarından büyük olamaz.")

            waste_disposition = "NONE"
            waste_type = ""
            waste_reason = ""
            waste_item_id = None
            waste_lot_no = ""
            waste_warehouse_id = None
            waste_setup_valid = True

            if float(waste_quantity) > 0.000001:
                waste_c1, waste_c2 = st.columns(2)
                waste_type = waste_c1.selectbox(
                    "Ziyan Türü",
                    WASTE_TYPES,
                    key=f"waste_type_{selected_order_id}",
                )
                disposition_options = ["RECOVERABLE", "DISPOSAL"]
                waste_disposition = waste_c2.selectbox(
                    "Ziyanın Değerlendirilmesi",
                    disposition_options,
                    format_func=lambda value: WASTE_DISPOSITION_LABELS[value],
                    key=f"waste_disposition_{selected_order_id}",
                )
                waste_reason = st.text_input(
                    "Ziyan Açıklaması (isteğe bağlı)",
                    placeholder="Örnek: Hat başlangıcında renk uygunluğu sağlanana kadar ayrılan ürün",
                    key=f"waste_reason_{selected_order_id}",
                )

                if waste_disposition == "RECOVERABLE":
                    k2_items = query_rows(
                        """
                        SELECT id, code, name, unit
                        FROM items
                        WHERE code = 'HM-K2' AND active = 1
                        """
                    )
                    waste_warehouses = active_warehouses()
                    if not k2_items or not waste_warehouses:
                        waste_setup_valid = False
                        st.error(
                            "Tekrar kullanılabilir ziyan için aktif K2 ERİTMELİK kartı ve depo bulunmalıdır."
                        )
                    else:
                        waste_item = k2_items[0]
                        waste_item_id = waste_item["id"]
                        st.info(
                            f"Bu miktar {waste_item['code']} | {waste_item['name']} stok kartına eklenecek."
                        )
                        waste_warehouse_labels = option_map(
                            waste_warehouses,
                            lambda row: f"{row['code']} | {row['name']}",
                        )
                        warehouse_ids = list(waste_warehouse_labels)
                        default_warehouse_index = (
                            warehouse_ids.index(order["output_warehouse_id"])
                            if order["output_warehouse_id"] in warehouse_ids
                            else 0
                        )
                        waste_wc1, waste_wc2 = st.columns(2)
                        waste_lot_no = waste_wc1.text_input(
                            "Ziyan / K2 LOT No",
                            value=f"K2-{order['product_lot_no']}",
                            key=f"waste_lot_{selected_order_id}",
                        )
                        waste_warehouse_id = waste_wc2.selectbox(
                            "Ziyanın Aktarılacağı Depo",
                            warehouse_ids,
                            index=default_warehouse_index,
                            format_func=lambda value: waste_warehouse_labels[value],
                            key=f"waste_warehouse_{selected_order_id}",
                        )
                else:
                    st.warning(
                        "Kullanılamayan ziyan stoklara eklenmeyecek; üretim emrinde ziyan kaydı olarak saklanacak."
                    )

            st.markdown("**Fiili Tüketimler**")
            actual_consumptions = {}
            for requirement in requirements:
                actual_consumptions[requirement["item_id"]] = st.number_input(
                    f"{requirement['name']} ({requirement['unit']})",
                    min_value=0.0,
                    value=float(requirement["planned_quantity"]),
                    step=0.001,
                    format="%.6f",
                    key=f"actual_consumption_{selected_order_id}_{requirement['item_id']}",
                )
            submitted = st.button(
                "Üretimi Tamamla ve Stokları İşle",
                type="primary",
                use_container_width=True,
                disabled=not waste_setup_valid,
                key=f"complete_order_button_{selected_order_id}",
            )
            if submitted:
                try:
                    complete_production_order(
                        selected_order_id,
                        gross_output,
                        actual_consumptions,
                        waste_quantity=waste_quantity,
                        waste_disposition=waste_disposition,
                        waste_type=waste_type,
                        waste_reason=waste_reason,
                        waste_item_id=waste_item_id,
                        waste_lot_no=waste_lot_no,
                        waste_warehouse_id=waste_warehouse_id,
                    )
                    success_text = (
                        f"✅ Üretim tamamlandı. Brüt: {float(gross_output):,.3f} "
                        f"{order['product_unit']} · Net mamul: {net_output:,.3f} "
                        f"{order['product_unit']} · Ziyan: {float(waste_quantity):,.3f} "
                        f"{order['product_unit']}."
                    )
                    if waste_disposition == "RECOVERABLE":
                        success_text += " Tekrar kullanılabilir ziyan K2 ERİTMELİK stoğuna işlendi."
                    elif waste_disposition == "DISPOSAL":
                        success_text += " Kullanılamayan ziyan stok oluşturmadan kaydedildi."
                    st.success(success_text)
                    st.rerun()
                except Exception as exc:
                    show_error(exc)

    with tab_recipes:
        recipe_editor()


def shipment_page():
    st.header("🚚 Müşteri Sevkiyatı")
    tab_new, tab_history = st.tabs(["🚚 Yeni Sevkiyat", "📋 Sevkiyat Geçmişi ve İptal"])

    with tab_new:
        available = stock_balance_df(positive_only=True, categories=["Mamul"])
        if available.empty:
            st.warning(
                "Sevk edilebilir Mamul stoku bulunmuyor. Üretilecek ürün için Mamul kategorisinde "
                "stok kartı ve reçete tanımlayıp üretimi tamamlayın."
            )
        else:
            row_options = {}
            for _, row in available.iterrows():
                label = (
                    f"{row['Malzeme / Ürün']} | LOT {row['LOT No']} | {row['Depo']} | "
                    f"{row['Güncel Stok']:,.3f} {row['Birim']}"
                )
                row_options[label] = row
            selected_label = st.selectbox("Sevk Edilecek Mamul", list(row_options), key="shipment_stock")
            selected = row_options[selected_label]
            with st.form("shipment_form", clear_on_submit=True):
                c1, c2 = st.columns(2)
                customer = c1.text_input("Müşteri Firma Adı")
                dispatch_no = c2.text_input("İrsaliye Numarası")
                plate = c1.text_input("Araç Plakası")
                quantity = c2.number_input(
                    f"Sevk Miktarı ({selected['Birim']})",
                    min_value=0.000001,
                    max_value=float(selected["Güncel Stok"]),
                    value=float(min(1000.0, selected["Güncel Stok"])),
                    step=1.0,
                    format="%.3f",
                )
                notes = st.text_area("Sevkiyat Notu")
                submitted = st.form_submit_button("Sevkiyatı Onayla", use_container_width=True)
            if submitted:
                try:
                    ref = create_shipment(
                        int(selected["item_id"]),
                        int(selected["lot_id"]),
                        int(selected["warehouse_id"]),
                        quantity,
                        customer,
                        dispatch_no,
                        plate,
                        notes,
                    )
                    st.success(f"✅ Sevkiyat tamamlandı. Sevkiyat no: {ref}")
                    st.rerun()
                except Exception as exc:
                    show_error(exc)

    with tab_history:
        shipments = shipments_df()
        if shipments.empty:
            st.info("Kayıtlı sevkiyat bulunmuyor.")
        else:
            display_frame(shipments, ["id"])
            active_shipments = shipments[shipments["Durum"] == "Tamamlandı"]
            if not active_shipments.empty:
                shipment_labels = {
                    int(row["id"]): f"{row['Sevkiyat No']} | {row['Müşteri']} | {row['Ürün']} | {row['Miktar']:,.3f} {row['Birim']}"
                    for _, row in active_shipments.iterrows()
                }
                selected_shipment_id = st.selectbox(
                    "İptal İşlemi Yapılacak Sevkiyat",
                    list(shipment_labels),
                    format_func=lambda value: shipment_labels[value],
                )
                if st.button("⚠️ Sevkiyat İptal İşlemini Aç", use_container_width=True):
                    st.session_state["confirm_shipment_cancel"] = selected_shipment_id
                if st.session_state.get("confirm_shipment_cancel") == selected_shipment_id:
                    reason = st.text_input("Sevkiyat İptal Nedeni", key=f"shipment_cancel_reason_{selected_shipment_id}")
                    c1, c2 = st.columns(2)
                    if c1.button("İptali Kesinleştir", key=f"confirm_shipment_{selected_shipment_id}", use_container_width=True):
                        try:
                            cancel_shipment(selected_shipment_id, reason)
                            st.session_state.pop("confirm_shipment_cancel", None)
                            st.success("✅ Sevkiyat iptal edildi ve mamul stoğu geri yüklendi.")
                            st.rerun()
                        except Exception as exc:
                            show_error(exc)
                    if c2.button("Vazgeç", key=f"abort_shipment_{selected_shipment_id}", use_container_width=True):
                        st.session_state.pop("confirm_shipment_cancel", None)
                        st.rerun()


def stock_cards_page():
    st.header("🗂️ Stok Kartı Oluşturma")
    st.info(
        "Kartları kategori alt başlıklarından oluşturabilir veya düzenleyebilirsiniz. "
        "Stok kodunu boş bırakırsanız sistem otomatik kod üretir."
    )
    category_tabs = st.tabs(
        [
            "🧱 Hammadde",
            "🧪 Yardımcı Kimyasallar",
            "📦 Ambalaj",
            "⚙️ Ara Mamul",
            "🏭 Mamul",
        ]
    )
    category_keys = {
        "Hammadde": "hammadde",
        "Yardımcı Kimyasal": "yardimci",
        "Ambalaj": "ambalaj",
        "Ara Mamul": "ara_mamul",
        "Mamul": "mamul",
    }
    for category, tab in zip(CATEGORIES, category_tabs):
        with tab:
            render_stock_card_category(category, category_keys[category])


def render_stock_card_category(category, key_prefix):
    cards = query_rows(
        """
        SELECT id, code, name, category, unit, min_stock, active
        FROM items WHERE category = ? ORDER BY active DESC, name
        """,
        (category,),
    )
    active_count = sum(1 for card in cards if card["active"])
    st.caption(f"{category} · {active_count} aktif kart")
    tab_new, tab_manage = st.tabs(["➕ Yeni Kart Oluştur", "✏️ Kartları Yönet"])

    with tab_new:
        with st.form(f"new_item_form_{key_prefix}", clear_on_submit=True):
            c1, c2 = st.columns(2)
            code = c1.text_input(
                "Stok Kodu (boşsa otomatik oluşur)",
                key=f"new_item_code_{key_prefix}",
            )
            name = c2.text_input("Stok Kartı Adı", key=f"new_item_name_{key_prefix}")
            unit = c1.selectbox("Birim", UNITS, key=f"new_item_unit_{key_prefix}")
            min_stock = c2.number_input(
                "Minimum Stok Seviyesi",
                min_value=0.0,
                value=0.0,
                step=1.0,
                key=f"new_item_min_{key_prefix}",
            )
            st.caption(f"Kart kategorisi otomatik olarak {category} olacaktır.")
            submitted = st.form_submit_button("Stok Kartını Oluştur", use_container_width=True)
        if submitted:
            try:
                created_code = create_item(code, name, category, unit, min_stock)
                st.success(f"✅ {name} kartı {created_code} koduyla oluşturuldu.")
                st.rerun()
            except Exception as exc:
                show_error(exc)

    with tab_manage:
        if not cards:
            st.info(f"{category} kategorisinde kayıtlı kart bulunmuyor.")
        else:
            card_labels = {
                card["id"]: (
                    f"{card['code']} | {card['name']} | "
                    f"{'Aktif' if card['active'] else 'Pasif'}"
                )
                for card in cards
            }
            item_id = st.selectbox(
                "Düzenlenecek Kart",
                list(card_labels),
                format_func=lambda value: card_labels[value],
                key=f"manage_item_select_{key_prefix}",
            )
            item = next(card for card in cards if card["id"] == item_id)
            with st.form(f"manage_item_form_{key_prefix}_{item_id}"):
                st.text_input("Stok Kodu", value=item["code"], disabled=True)
                name = st.text_input(
                    "Stok Kartı Adı",
                    value=item["name"],
                    key=f"manage_item_name_{key_prefix}_{item_id}",
                )
                c1, c2 = st.columns(2)
                unit = c1.selectbox(
                    "Birim",
                    UNITS,
                    index=UNITS.index(item["unit"]),
                    key=f"manage_item_unit_{key_prefix}_{item_id}",
                )
                min_stock = c2.number_input(
                    "Minimum Stok Seviyesi",
                    min_value=0.0,
                    value=float(item["min_stock"]),
                    step=1.0,
                    key=f"manage_item_min_{key_prefix}_{item_id}",
                )
                active = st.checkbox(
                    "Aktif Kart",
                    value=bool(item["active"]),
                    key=f"manage_item_active_{key_prefix}_{item_id}",
                )
                submitted = st.form_submit_button("Değişiklikleri Kaydet", use_container_width=True)
            if submitted:
                try:
                    update_item(item_id, name, category, unit, min_stock, active)
                    st.success("✅ Stok kartı güncellendi; geçmiş hareket bağlantıları korundu.")
                    st.rerun()
                except Exception as exc:
                    show_error(exc)

        cards_df = query_df(
            """
            SELECT code AS "Stok Kodu", name AS "Stok Kartı", unit AS "Birim",
                   min_stock AS "Minimum Stok",
                   CASE WHEN active = 1 THEN 'Aktif' ELSE 'Pasif' END AS "Durum"
            FROM items WHERE category = ? ORDER BY active DESC, name
            """,
            (category,),
        )
        display_frame(cards_df)


def definitions_and_reports_page():
    st.header("📈 Raporlar ve Tanımlar")
    tab_reports, tab_warehouses, tab_system = st.tabs(
        ["📊 Raporlar", "🏬 Depolar", "⚙️ Sistem"]
    )

    with tab_reports:
        c1, c2 = st.columns(2)
        start_date = c1.date_input("Rapor Başlangıç Tarihi", value=date(date.today().year, 1, 1))
        end_date = c2.date_input("Rapor Bitiş Tarihi", value=date.today())
        if start_date > end_date:
            st.error("Başlangıç tarihi bitiş tarihinden sonra olamaz.")
        else:
            movements = movement_history_df(start_date, end_date)
            c3, c4, c5 = st.columns(3)
            c3.metric("Dönem Hareket Sayısı", len(movements))
            c4.metric("Üretim Emri Sayısı", len(production_orders_df()))
            c5.metric("Sevkiyat Kaydı", len(shipments_df(start_date, end_date)))
            if not movements.empty:
                display_frame(movements)
            report_bytes = build_excel_report(start_date, end_date)
            st.download_button(
                "📥 Kurumsal Excel Raporunu İndir",
                data=report_bytes,
                file_name=f"PET_ERP_Rapor_{start_date}_{end_date}.xlsx",
                mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                use_container_width=True,
            )
            st.caption("Anlık Stok sayfası güncel bakiyeyi; diğer sayfalar seçilen dönem hareketlerini gösterir.")

    with tab_warehouses:
        with st.form("new_warehouse_form", clear_on_submit=True):
            c1, c2 = st.columns(2)
            code = c1.text_input("Yeni Depo Kodu")
            name = c2.text_input("Yeni Depo Adı")
            warehouse_type = st.selectbox("Depo Türü", ["Genel", "Hammadde", "Mamul", "Ambalaj", "Silo"])
            submitted = st.form_submit_button("Depoyu Kaydet", use_container_width=True)
        if submitted:
            try:
                create_warehouse(code, name, warehouse_type)
                st.success("✅ Depo tanımlandı.")
                st.rerun()
            except Exception as exc:
                show_error(exc)

        all_warehouses = query_rows(
            "SELECT id, code, name, warehouse_type, active FROM warehouses ORDER BY name"
        )
        if all_warehouses:
            warehouse_labels = option_map(
                all_warehouses,
                lambda row: f"{row['code']} | {row['name']} | {'Aktif' if row['active'] else 'Pasif'}",
            )
            selected_id = st.selectbox(
                "Durumu Değiştirilecek Depo",
                list(warehouse_labels),
                format_func=lambda value: warehouse_labels[value],
            )
            selected = next(row for row in all_warehouses if row["id"] == selected_id)
            desired_active = st.checkbox("Depo Aktif", value=bool(selected["active"]), key=f"warehouse_active_{selected_id}")
            if st.button("Depo Durumunu Güncelle", use_container_width=True):
                try:
                    set_warehouse_active(selected_id, desired_active)
                    st.success("✅ Depo durumu güncellendi.")
                    st.rerun()
                except Exception as exc:
                    show_error(exc)
        warehouses_df = query_df(
            """
            SELECT code AS "Depo Kodu", name AS "Depo Adı", warehouse_type AS "Depo Türü",
                   CASE WHEN active = 1 THEN 'Aktif' ELSE 'Pasif' END AS "Durum"
            FROM warehouses ORDER BY name
            """
        )
        display_frame(warehouses_df)

    with tab_system:
        st.subheader("Veri Güvenliği")
        st.write(
            "Tüm operasyonel kayıtlar SQLite veritabanında kalıcı olarak tutulur. "
            "İptal işlemleri geçmiş hareketleri silmez; ters kayıt oluşturur."
        )
        backup = build_database_backup()
        st.download_button(
            "💾 Veritabanı Yedeğini İndir",
            data=backup,
            file_name=f"PET_ERP_Yedek_{datetime.now().strftime('%Y%m%d_%H%M')}.db",
            mime="application/octet-stream",
            use_container_width=True,
        )
        st.code("streamlit run PET_Resin_ERP_v2_7.py", language="bash")
        st.caption(f"Veritabanı dosyası uygulama klasöründe oluşturulur: {DB_PATH.name}")


def main():
    st.set_page_config(page_title=f"{APP_NAME} {APP_VERSION}", layout="wide")
    st.markdown(
        """
        <style>
        .block-container {padding-top: 1.2rem; padding-bottom: 2rem;}
        [data-testid="stMetric"] {background: #f7f9fc; border: 1px solid #dfe6ef; padding: 14px; border-radius: 10px;}
        .stButton > button, .stDownloadButton > button {border-radius: 8px;}
        </style>
        """,
        unsafe_allow_html=True,
    )
    try:
        initialize_database()
    except Exception as exc:
        st.error(f"Veritabanı başlatılamadı: {exc}")
        st.stop()

    st.sidebar.title("🧪 PET Resin ERP")
    st.sidebar.caption(f"{APP_VERSION} · Kalıcı ve lot izlenebilir")
    page = st.sidebar.radio(
        "Ana Menü",
        [
            "📊 Ana Panel",
            "🗂️ Stok Kartı Oluşturma",
            "📦 Stok Yönetimi",
            "🏭 Üretim",
            "🚚 Sevkiyat",
            "📈 Raporlar ve Tanımlar",
        ],
    )
    st.sidebar.divider()
    st.sidebar.success("SQLite bağlantısı aktif")

    if page == "📊 Ana Panel":
        dashboard_page()
    elif page == "🗂️ Stok Kartı Oluşturma":
        stock_cards_page()
    elif page == "📦 Stok Yönetimi":
        stock_page()
    elif page == "🏭 Üretim":
        production_page()
    elif page == "🚚 Sevkiyat":
        shipment_page()
    else:
        definitions_and_reports_page()


if __name__ == "__main__":
    main()
